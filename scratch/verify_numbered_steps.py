import openpyxl
import re

wb_curr = openpyxl.load_workbook(r'd:\Project\BTN\SIT\Test Case.xlsx', data_only=True)
ws_curr = wb_curr['TC BTN SMART Web']

wb_bak = openpyxl.load_workbook(r'd:\Project\BTN\SIT\Test Case_backup_before_delete_redundant.xlsx', data_only=True)
ws_bak = wb_bak['TC BTN SMART Web']

def normalize(text):
    if not text:
        return ''
    t = str(text).lower()
    t = re.sub(r'[^a-z0-9]', '', t)
    return t

bak_items = []
for r in range(2, 1030):
    m = ws_bak.cell(r, 2).value
    s = ws_bak.cell(r, 3).value
    t = ws_bak.cell(r, 4).value
    steps = ws_bak.cell(r, 7).value
    exp = ws_bak.cell(r, 8).value
    if m and t:
        bak_items.append({
            'row': r,
            'mod': str(m).strip(),
            'sub': str(s).strip() if s else '',
            'title': str(t).strip(),
            'norm_mod': normalize(m),
            'norm_sub': normalize(s),
            'norm_title': normalize(t),
            'steps': str(steps or '').strip(),
            'exp': str(exp or '').strip()
        })

# Manual overrides for the 3 special cases
overrides = {
    377: (
        '1. Input nama prospek pada kolom pencarian.',
        '1. Menampilkan data sesuai kata kunci yang diinput.'
    ),
    380: (
        "1. Lakukan aktifitas marketing call dengan status call 'Menolak'.\n2. Lakukan aktifitas marketing call kembali dengan status lain.",
        "1. Berhasil melakukan call dengan status menolak.\n2. Menampilkan popup alert 'Anda tidak dapat memindahkan dari stage Ditolak ke Call'."
    ),
    382: (
        "1. Klik Input Data NOA Open Account.\n2. Input nomor rekening.\n3. Pilih tanggal dana masuk.\n4. Pilih apakah nasabah sudah mengaktifkan Mbanking.\n5. Input remark.\n6. Klik button Simpan.",
        "1. Menampilkan halaman input data NOA Open Account.\n2. Berhasil menginput nomor rekening.\n3. Berhasil memilih tanggal dana masuk.\n4. Berhasil memilih apakah nasabah sudah mengaktifkan mbanking atau belum.\n5. Berhasil menginput remark.\n6. Berhasil menyimpan data NOA Open Account, menampilkan popup 'berhasil melakukan input data open account sukses' dan status prospek atau staging berubah menjadi Pembukaan."
    )
}

matched_results = []
for r in range(2, ws_curr.max_row + 1):
    no = ws_curr.cell(r, 1).value
    m = ws_curr.cell(r, 2).value
    s = ws_curr.cell(r, 3).value
    t = ws_curr.cell(r, 4).value
    
    if r in overrides:
        matched_results.append((r, no, overrides[r][0], overrides[r][1]))
        continue
        
    nm = normalize(m)
    ns = normalize(s)
    nt = normalize(t)
    
    found = None
    for b in bak_items:
        if b['norm_mod'] == nm and b['norm_sub'] == ns and b['norm_title'] == nt:
            found = b
            break
            
    if not found:
        for b in bak_items:
            if b['norm_mod'] == nm and b['norm_title'] == nt:
                found = b
                break
                
    if not found:
        for b in bak_items:
            if b['norm_mod'] == nm and (nt in b['norm_title'] or b['norm_title'] in nt):
                found = b
                break
                
    if found:
        matched_results.append((r, no, found['steps'], found['exp']))
    else:
        print(f"FAILED TO MATCH: row {r}, {no}, {m}, {s}, {t}")

print(f"Total processed: {len(matched_results)} / {ws_curr.max_row - 1}")

# Verify that all 921 have numbered steps
has_numbered_steps = 0
for r, no, st, ex in matched_results:
    if re.search(r'^\s*1[\.\)]', st, re.M):
        has_numbered_steps += 1
    else:
        print(f"Row {r} ({no}) step does NOT start with '1.': {repr(st)[:60]}")

print(f"Total with '1.' numbered steps: {has_numbered_steps} / {len(matched_results)}")
