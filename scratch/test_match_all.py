import openpyxl
import re

wb_curr = openpyxl.load_workbook(r'd:\Project\BTN\SIT\Test Case.xlsx', data_only=True)
ws_curr = wb_curr['TC BTN SMART Web']

wb_bak = openpyxl.load_workbook(r'd:\Project\BTN\SIT\Test Case_backup_before_delete_redundant.xlsx', data_only=True)
ws_bak = wb_bak['TC BTN SMART Web']

def normalize(text):
    if not text:
        return ''
    t = str(text).lower()
    t = re.sub(r'[^a-z0-9]', '', t)
    return t

# Build backup lookup
bak_items = []
for r in range(2, 1030):
    m = ws_bak.cell(r, 2).value
    s = ws_bak.cell(r, 3).value
    t = ws_bak.cell(r, 4).value
    steps = ws_bak.cell(r, 7).value
    exp = ws_bak.cell(r, 8).value
    if m and t:
        bak_items.append({
            'row': r,
            'mod': str(m).strip(),
            'sub': str(s).strip() if s else '',
            'title': str(t).strip(),
            'norm_mod': normalize(m),
            'norm_sub': normalize(s),
            'norm_title': normalize(t),
            'steps': steps,
            'exp': exp
        })

print(f"Total backup items: {len(bak_items)}")

# Match each current row
matched_count = 0
unmatched_rows = []

for r in range(2, ws_curr.max_row + 1):
    no = ws_curr.cell(r, 1).value
    m = ws_curr.cell(r, 2).value
    s = ws_curr.cell(r, 3).value
    t = ws_curr.cell(r, 4).value
    
    nm = normalize(m)
    ns = normalize(s)
    nt = normalize(t)
    
    # 1. Exact match (mod, sub, title)
    found = None
    for b in bak_items:
        if b['norm_mod'] == nm and b['norm_sub'] == ns and b['norm_title'] == nt:
            found = b
            break
            
    # 2. Match without submenu if sub differs slightly
    if not found:
        for b in bak_items:
            if b['norm_mod'] == nm and b['norm_title'] == nt:
                found = b
                break
                
    # 3. Match substring / fuzzy in same module
    if not found:
        for b in bak_items:
            if b['norm_mod'] == nm and (nt in b['norm_title'] or b['norm_title'] in nt):
                found = b
                break

    if found:
        matched_count += 1
    else:
        unmatched_rows.append((r, no, m, s, t))

print(f"Matched: {matched_count} / {ws_curr.max_row - 1}")
print(f"Unmatched: {len(unmatched_rows)}")
for u in unmatched_rows:
    print(" ", u)
