import os
import openpyxl

EXCEL_PATH = r'd:\Project\BTN\SIT\Test Case.xlsx'
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['TC BTN SMART Web']

print(f"Total rows in sheet: {ws.max_row}")

# Check which test cases have screenshots in d:\Project\BTN\Hasil Uji\Screenshot\Web
ss_base = r'd:\Project\BTN\Hasil Uji\Screenshot\Web'
found_imgs = []
for r in range(2, ws.max_row + 1):
    no = str(ws.cell(r, 1).value).strip()
    sec_num = int(no.split('.')[0])
    m = str(ws.cell(r, 2).value).strip()
    s = str(ws.cell(r, 3).value or '').strip()
    t = str(ws.cell(r, 4).value).strip()
    
    sub_label = f"{m} - {s}" if s else m
    folder_sub = f"{sec_num:02d}. {sub_label}"
    folder_tc = f"{no} {t}"
    
    tc_dir = os.path.join(ss_base, folder_sub, folder_tc)
    if os.path.exists(tc_dir):
        imgs = [f for f in os.listdir(tc_dir) if f.endswith(('.png', '.jpg', '.jpeg'))]
        if imgs:
            found_imgs.append((no, t, imgs))

print(f"Total TCs with images found: {len(found_imgs)}")
for no, t, imgs in found_imgs:
    print(f"  {no} {t}: {imgs}")
