import docx
import openpyxl

wb = openpyxl.load_workbook(r'd:\Project\BTN\SIT\Test Case.xlsx', data_only=True)
ws = wb['TC BTN SMART Web']

# Map (module, title_clean) -> sub_menu
excel_map = {}
for r in range(2, ws.max_row + 1):
    m = ws.cell(r, 2).value
    s = ws.cell(r, 3).value
    t = ws.cell(r, 4).value
    if m and t:
        m_str = str(m).strip()
        s_str = str(s).strip() if s else ''
        t_clean = str(t).strip().lower()
        excel_map[(m_str, t_clean)] = s_str

doc = docx.Document(r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web (Updated).docx')
body = doc.element.body

modules_list = [
    'Login', 'Profile Nasabah & Sales', 'Profile', 'Overview', 'User Authority', 
    'Bisnis dan Produk', 'Kantor', 'Menu Target', 'Upload Bulk', 
    'Lead Generation', 'Lead Qualification', 'List Prospek ETB', 
    'Menu Absent', 'Sales Tracking Activity', 'Setting Absent', 
    'Sales Force', 'Report Funding', 'Report Lending', 
    'Re-Assign dan Approval', 'Export data management'
]

cur_mod = 'Unknown'
extracted = []

for child in list(body):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        t = ''.join(child.itertext()).strip()
        if t and 'Modul' in t:
            for m in modules_list:
                if m in t:
                    cur_mod = m
                    break
    elif tag == 'tbl':
        rows = child.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
        if rows:
            cells = rows[0].findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
            if len(cells) == 7 and len(rows) > 1:
                for r in rows[1:]:
                    r_cells = r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
                    no = ''.join(r_cells[0].itertext()).strip()
                    title = ''.join(r_cells[1].itertext()).strip()
                    steps = ''.join(r_cells[2].itertext()).strip()
                    exp = ''.join(r_cells[3].itertext()).strip()
                    act = ''.join(r_cells[4].itertext()).strip()
                    st = ''.join(r_cells[5].itertext()).strip()
                    rem = ''.join(r_cells[6].itertext()).strip()
                    
                    # Look up sub menu from excel_map
                    sub = excel_map.get((cur_mod, title.lower()))
                    if sub is None:
                        # try partial match within same module
                        for (em, et), es in excel_map.items():
                            if em == cur_mod and (et in title.lower() or title.lower() in et):
                                sub = es
                                break
                    if sub is None:
                        sub = ''
                    extracted.append({
                        'module': cur_mod,
                        'submenu': sub,
                        'no': no,
                        'title': title,
                        'steps': steps,
                        'expected': exp,
                        'actual': act,
                        'status': st,
                        'remarks': rem
                    })

print(f"Total extracted: {len(extracted)}")
unmatched = [x for x in extracted if x['submenu'] == '' and x['module'] not in ['Login', 'Profile', 'Menu Target', 'List Prospek ETB']]
print(f"Unmatched submenus in modules that have submenus: {len(unmatched)}")
if unmatched:
    for u in unmatched[:10]:
        print(f"  Modul: {u['module']} | Title: {u['title']}")
