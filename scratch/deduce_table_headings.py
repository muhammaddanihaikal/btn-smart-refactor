import docx
import openpyxl

# Load Excel to map TC titles to (Module, Sub Menu)
wb = openpyxl.load_workbook(r'd:\Project\BTN\SIT\Test Case.xlsx', data_only=True)
ws = wb['TC BTN SMART Web']

excel_tcs = {}
for r in range(2, ws.max_row + 1):
    mod = ws.cell(r, 2).value
    sub = ws.cell(r, 3).value
    title = ws.cell(r, 4).value
    if title:
        excel_tcs[str(title).strip().lower()] = (
            str(mod).strip() if mod else '',
            str(sub).strip() if sub else ''
        )

doc = docx.Document(r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web (Updated).docx')
data_tables = [t for t in doc.tables if len(t.columns) == 7 and len(t.rows) > 1]

print(f"Total tables: {len(data_tables)}")

for idx, t in enumerate(data_tables):
    # Check all titles in this table to find the best matching (Module, Sub Menu)
    matched_subs = []
    matched_mods = []
    
    titles_in_table = []
    for r in t.rows[1:]:
        title = r.cells[1].text.strip()
        titles_in_table.append(title)
        t_clean = title.lower()
        if t_clean in excel_tcs:
            m, s = excel_tcs[t_clean]
            matched_mods.append(m)
            matched_subs.append(s)
        else:
            # Substring match
            for k, (m, s) in excel_tcs.items():
                if k in t_clean or t_clean in k:
                    matched_mods.append(m)
                    matched_subs.append(s)
                    break
                    
    # Most common mod and sub
    from collections import Counter
    mod = Counter(matched_mods).most_common(1)[0][0] if matched_mods else "Unknown"
    sub = Counter(matched_subs).most_common(1)[0][0] if matched_subs else ""
    
    label = f"Modul {mod} - {sub}" if sub else f"Modul {mod}"
    print(f"Table {idx+1:2d} ({len(t.rows)-1:2d} TCs) -> '{label}' | First TC: {titles_in_table[0][:30]}")
