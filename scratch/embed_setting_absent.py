import os
import re
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def win_p(p):
    p = os.path.abspath(p)
    return '\\\\?\\' + p if not p.startswith('\\\\?\\') else p

doc_p = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Web.docx'
screenshots_base = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Web'

doc = docx.Document(doc_p)

def natural_sort_key(s):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

def get_images_for_tc(module_name, tc_str):
    mod_dir = os.path.join(screenshots_base, module_name)
    for d in os.listdir(win_p(mod_dir)):
        if d.startswith(tc_str + ' '):
            fld = os.path.join(mod_dir, d)
            if os.path.isdir(win_p(fld)):
                imgs = [f for f in os.listdir(win_p(fld)) if f.lower().endswith(('.png', '.jpg', '.jpeg'))]
                imgs.sort(key=natural_sort_key)
                return [os.path.join(fld, img) for img in imgs]
    return []

def format_expected_result_p(p, custom_desc=None):
    if custom_desc is None:
        text = p.text.strip()
        if not text.startswith('Hasil yang diharapkan'):
            return
        parts = text.split('[STATUS]:')
        if len(parts) < 2:
            return
        desc_full = parts[1].strip().lstrip('\r\n ')
        if '[Success]' in desc_full:
            desc = re.sub(r'\s*-\s*\[Success\]', '', desc_full).strip()
        else:
            desc = desc_full
    else:
        desc = custom_desc.strip()
        if '[Success]' in desc:
            desc = re.sub(r'\s*-\s*\[Success\]', '', desc).strip()
            
    p.text = ''
    
    # Run 1: Header
    r1 = p.add_run('Hasil yang diharapkan [STATUS]:')
    r1.font.name = 'Arial'
    r1.font.size = Pt(10)
    r1.bold = True
    
    # Run 2: Break
    r_br = p.add_run()
    r_br.add_break()
    
    # Run 3: Description
    r2 = p.add_run(desc)
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    rPr2 = r2._element.get_or_add_rPr()
    c2 = OxmlElement('w:color')
    c2.set(qn('w:val'), '2C5293')
    rPr2.append(c2)
    
    # Run 4: - [Success]
    r3 = p.add_run(' - [Success]')
    r3.font.name = 'Arial'
    r3.font.size = Pt(10)
    r3.bold = True
    rPr3 = r3._element.get_or_add_rPr()
    c3 = OxmlElement('w:color')
    c3.set(qn('w:val'), '2C5293')
    rPr3.append(c3)

# Find tables for Modul 40, 41, 42
mod_tables = {}
current_h = ''

for el in doc.element.body:
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(el, doc)
        if p.style.name.startswith('Heading 1'):
            current_h = p.text
    elif tag == 'tbl':
        for m in ['40', '41', '42']:
            if f"{m}. Modul Setting Absent" in current_h:
                if m not in mod_tables:
                    mod_tables[m] = []
                mod_tables[m].append(docx.table.Table(el, doc))

for m in ['40', '41', '42']:
    tbls = mod_tables.get(m, [])
    print(f"Modul {m} tables: {len(tbls)} (Table 0: {len(tbls[0].rows)} rows, Table 1: {len(tbls[1].rows)} rows)")

print("\n=== 1. EMBEDDING MODUL 40 (Attendance Spot) ===")
tbl40_0, tbl40_1 = mod_tables['40'][0], mod_tables['40'][1]

# 40.1 in tbl40_0 (row 2 image, row 3 expected)
imgs_40_1 = get_images_for_tc('40. Setting Absent - Attendance Spot', '40.1')
c40_0 = tbl40_0.rows[2].cells[1]
c40_0.text = ''
p40_0 = c40_0.paragraphs[0]
p40_0.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, img in enumerate(imgs_40_1):
    p = p40_0 if i == 0 else c40_0.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img, width=Inches(5.9))
format_expected_result_p(tbl40_0.rows[3].cells[1].paragraphs[0])
print(f"Embedded 40.1: {len(imgs_40_1)} images")

# 40.2 to 40.19 in tbl40_1
for k in range(2, 20):
    tc_str = f"40.{k}"
    idx = k - 2
    r_img = idx * 3 + 1
    r_exp = idx * 3 + 2
    
    c = tbl40_1.rows[r_img].cells[1]
    c.text = ''
    p_img = c.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    imgs = get_images_for_tc('40. Setting Absent - Attendance Spot', tc_str)
    for i, img in enumerate(imgs):
        p = p_img if i == 0 else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(5.9))
        
    format_expected_result_p(tbl40_1.rows[r_exp].cells[1].paragraphs[0])
    print(f"Embedded {tc_str}: {len(imgs)} images")

print("\n=== 2. EMBEDDING MODUL 41 (Work Pattern) ===")
tbl41_0, tbl41_1 = mod_tables['41'][0], mod_tables['41'][1]

# 41.1 in tbl41_0
imgs_41_1 = get_images_for_tc('41. Setting Absent - Work Pattern', '41.1')
c41_0 = tbl41_0.rows[2].cells[1]
c41_0.text = ''
p41_0 = c41_0.paragraphs[0]
p41_0.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, img in enumerate(imgs_41_1):
    p = p41_0 if i == 0 else c41_0.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img, width=Inches(5.9))
format_expected_result_p(tbl41_0.rows[3].cells[1].paragraphs[0])
print(f"Embedded 41.1: {len(imgs_41_1)} images")

# 41.2 to 41.19 in tbl41_1
for k in range(2, 20):
    tc_str = f"41.{k}"
    idx = k - 2
    r_img = idx * 3 + 1
    r_exp = idx * 3 + 2
    
    c = tbl41_1.rows[r_img].cells[1]
    c.text = ''
    p_img = c.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    imgs = get_images_for_tc('41. Setting Absent - Work Pattern', tc_str)
    for i, img in enumerate(imgs):
        p = p_img if i == 0 else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(5.9))
        
    format_expected_result_p(tbl41_1.rows[r_exp].cells[1].paragraphs[0])
    print(f"Embedded {tc_str}: {len(imgs)} images")

print("\n=== 3. EMBEDDING MODUL 42 (Holiday) ===")
tbl42_0, tbl42_1 = mod_tables['42'][0], mod_tables['42'][1]

# Expected results mapping for 42:
exp_results_42 = {
    '42.1': 'Berhasil menampilkan halaman Holiday.',
    '42.2': 'Berhasil menampilkan data sesuai keyword.',
    '42.3': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '42.4': 'Berhasil menyimpan data Holiday.',
    '42.5': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '42.6': 'Berhasil menyimpan perubahan data Holiday.',
    '42.7': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '42.8': 'Berhasil menghapus data Holiday.',
}

# 42.1 in tbl42_0
imgs_42_1 = get_images_for_tc('42. Setting Absent - Holiday', '42.1')
c42_0 = tbl42_0.rows[2].cells[1]
c42_0.text = ''
p42_0 = c42_0.paragraphs[0]
p42_0.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, img in enumerate(imgs_42_1):
    p = p42_0 if i == 0 else c42_0.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img, width=Inches(5.9))
format_expected_result_p(tbl42_0.rows[3].cells[1].paragraphs[0], exp_results_42['42.1'])
print(f"Embedded 42.1: {len(imgs_42_1)} images")

# 42.2 to 42.8 in tbl42_1
for k in range(2, 9):
    tc_str = f"42.{k}"
    idx = k - 2
    r_img = idx * 3 + 1
    r_exp = idx * 3 + 2
    
    c = tbl42_1.rows[r_img].cells[1]
    c.text = ''
    p_img = c.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    imgs = get_images_for_tc('42. Setting Absent - Holiday', tc_str)
    for i, img in enumerate(imgs):
        p = p_img if i == 0 else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(5.9))
        
    format_expected_result_p(tbl42_1.rows[r_exp].cells[1].paragraphs[0], exp_results_42[tc_str])
    print(f"Embedded {tc_str}: {len(imgs)} images | Expected: {exp_results_42[tc_str]}")

doc.save(doc_p)
print(f"\nSuccessfully saved {doc_p} with all Modul 40, 41, 42 cropped images embedded & formatted!")
