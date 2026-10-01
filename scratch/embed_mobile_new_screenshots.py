import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc_path = r'D:/Project/BTN Smart/Refactor/Hasil Uji/Dokumen_Hasil_Uji_Mobile.docx'
doc = docx.Document(doc_path)

# 1. Embed Table 59 (3.2)
t59 = doc.tables[59]
c59 = t59.rows[1].cells[1]
p59 = c59.paragraphs[0]
p59.text = ''
p59.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld_3_2 = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile\03. Sales Tracking Activity - Visit In\3.2 Melakukan Visit in menggunakan sales Funding'
p_img1 = os.path.join(fld_3_2, '1.png')
p_img2 = os.path.join(fld_3_2, '2.jpeg')
r = p59.add_run()
r.add_picture(p_img1, height=Inches(3.8))
p59.add_run('  ')
r = p59.add_run()
r.add_picture(p_img2, height=Inches(3.8))
print('Embedded 3.2')

# Modul 10 base folder
fld_10_base = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile\10. Prospek & Nasabah - Aktifitas Prospek'

# Helper to embed in a table
def embed_table(table, folder_name, img_names):
    c = table.rows[1].cells[1]
    p = c.paragraphs[0]
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = os.path.join(fld_10_base, folder_name)
    for i, img in enumerate(img_names):
        if i > 0:
            p.add_run('  ')
        r = p.add_run()
        r.add_picture(os.path.join(fld, img), height=Inches(3.8))

# 2. Table 132 (10.20)
embed_table(doc.tables[132], '10.20 Melakukan Submit NOA Open Account tanpa field mandatory', ['01.png'])
print('Embedded 10.20')

# 3. Table 133 (10.21)
embed_table(doc.tables[133], '10.21 Membuka tab Top Up data pada halaman aktifitas Marketing prospek ETB', ['01.png'])
print('Embedded 10.21')

# 4. Table 134 (10.22)
embed_table(doc.tables[134], '10.22 Melakukan Submit Top Up dana dengan data valid', ['01.png', '02.png'])
print('Embedded 10.22')

# 5. Table 135 (10.23)
embed_table(doc.tables[135], '10.23 Melakukan Submit Top Up dana tanpa field mandatory', ['01.png'])
print('Embedded 10.23')

doc.save(doc_path)
print('Successfully saved Dokumen_Hasil_Uji_Mobile.docx with new screenshots!')
