import docx
import openpyxl
import re
from collections import OrderedDict

EXCEL_PATH = r'd:\Project\BTN\SIT\Test Case.xlsx'
DOC_PATH = r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web (Updated).docx'

# 1. Load Excel mapping
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['TC BTN SMART Web']

excel_tc_map = {}
excel_sub_order = []
cur_pair = None

for r in range(2, ws.max_row + 1):
    m = ws.cell(r, 2).value
    s = ws.cell(r, 3).value
    t = ws.cell(r, 4).value
    if m and t:
        m_str = str(m).strip()
        s_str = str(s).strip() if s else ''
        t_clean = str(t).strip().lower()
        excel_tc_map[(m_str, t_clean)] = s_str
        
        pair = (m_str, s_str)
        if pair != cur_pair:
            cur_pair = pair
            excel_sub_order.append(pair)

# 2. Load docx and extract all TCs preserving Module context
doc = docx.Document(DOC_PATH)
body = doc.element.body

modules_list = [
    'Login', 'Profile Nasabah & Sales', 'Profile', 'Overview', 'User Authority', 
    'Bisnis dan Produk', 'Kantor', 'Menu Target', 'Upload Bulk', 
    'Lead Generation', 'Lead Qualification', 'List Prospek ETB', 
    'Menu Absent', 'Sales Tracking Activity', 'Setting Absent', 
    'Sales Force', 'Report Funding', 'Report Lending', 
    'Re-Assign dan Approval', 'Export data management'
]

cur_mod = 'Unknown'
raw_tcs = []

for child in list(body):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        t = ''.join(child.itertext()).strip()
        if t and 'Modul' in t:
            for m in modules_list:
                if m in t:
                    cur_mod = m
                    break
    elif tag == 'tbl':
        rows = child.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
        if rows:
            cells = rows[0].findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
            if len(cells) == 7 and len(rows) > 1:
                for r in rows[1:]:
                    r_cells = r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
                    no = ''.join(r_cells[0].itertext()).strip()
                    title = ''.join(r_cells[1].itertext()).strip()
                    steps = ''.join(r_cells[2].itertext()).strip()
                    exp = ''.join(r_cells[3].itertext()).strip()
                    act = ''.join(r_cells[4].itertext()).strip()
                    st = ''.join(r_cells[5].itertext()).strip()
                    rem = ''.join(r_cells[6].itertext()).strip()
                    
                    # Match sub menu
                    sub = excel_tc_map.get((cur_mod, title.lower()))
                    if sub is None:
                        for (em, et), es in excel_tc_map.items():
                            if em == cur_mod and (et in title.lower() or title.lower() in et):
                                sub = es
                                break
                    if sub is None:
                        sub = ''
                        
                    raw_tcs.append({
                        'module': cur_mod,
                        'submenu': sub,
                        'no': no,
                        'title': title,
                        'steps': steps,
                        'expected': exp,
                        'actual': act,
                        'status': st,
                        'remarks': rem
                    })

print(f"Total raw TCs extracted: {len(raw_tcs)}")

# 3. Filter out redundant "Menampilkan daftar" TCs
def is_redundant_list_tc(tc):
    t = tc['title'].strip().lower()
    # List of exact/prefix patterns to drop
    drop_prefixes = [
        'menampilkan daftar data',
        'menampilkan daftar permintaan',
        'menampilkan daftar perangkat',
        'menampilkan daftar log aktivitas',
        'menampilkan daftar ',
        'melihat daftar '
    ]
    # Keep specific non-redundant ones
    if 'tidak terdaftar' in t or 'di daftarkan' in t:
        return False
    if 'tab ' in t and ('rekening' in t or 'pipeline' in t):
        return False
    return any(t.startswith(p) for p in drop_prefixes)

filtered_tcs = [tc for tc in raw_tcs if not is_redundant_list_tc(tc)]
dropped_tcs = [tc for tc in raw_tcs if is_redundant_list_tc(tc)]

print(f"Dropped TCs count: {len(dropped_tcs)}")
for d in dropped_tcs:
    print(f"  Dropped: [{d['module']}] {d['no']} - {d['title']}")
print(f"Remaining TCs count: {len(filtered_tcs)}")

# 4. Group by (module, submenu)
grouped = OrderedDict()
for tc in filtered_tcs:
    key = (tc['module'], tc['submenu'])
    if key not in grouped:
        grouped[key] = []
    grouped[key].append(tc)

print(f"\nTotal Sub Menu groups formed: {len(grouped)}")
for (m, s), tcs in list(grouped.items())[:15]:
    label = f"Modul {m} - {s}" if s else f"Modul {m}"
    print(f"  {label:<45}: {len(tcs)} TCs")
