import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc_path = r'D:/Project/BTN Smart/Refactor/Hasil Uji/Dokumen_Hasil_Uji_Mobile.docx'
doc = docx.Document(doc_path)

t86 = doc.tables[86] # 7.3
t87 = doc.tables[87] # 7.4

def set_blocker_note(table, tc_type):
    # Row 1, Cell 1 -> Note content
    c1 = table.rows[1].cells[1]
    c1.text = ''
    
    # Optional light background
    shading_elm = parse_xml(r'<w:shd {} w:fill="FFFBEB"/>'.format(nsdecls('w')))
    c1._tc.get_or_add_tcPr().append(shading_elm)
    
    # Paragraph 1: Header
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p1.add_run('[STATUS: PENDING / ON HOLD - BLOCKER ISSUE]')
    r1.font.name = 'Arial'
    r1.font.size = Pt(10)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(185, 28, 28) # Dark Red
    
    # Paragraph 2: Detail Bug
    p2 = c1.add_paragraph()
    r2_b = p2.add_run('Deskripsi Masalah: ')
    r2_b.font.name = 'Arial'
    r2_b.font.size = Pt(9.5)
    r2_b.font.bold = True
    
    r2_t = p2.add_run(f'Komponen peta (Map) pada halaman detail log absensi {tc_type} tidak muncul (blank / gagal render) saat tester membuka riwayat detail absensi.')
    r2_t.font.name = 'Arial'
    r2_t.font.size = Pt(9.5)
    
    # Paragraph 3: Dampak & Tindakan
    p3 = c1.add_paragraph()
    r3_b = p3.add_run('Dampak & Tindakan: ')
    r3_b.font.name = 'Arial'
    r3_b.font.size = Pt(9.5)
    r3_b.font.bold = True
    
    r3_t = p3.add_run('Bukti screenshot pengujian ditunda sementara menunggu perbaikan render map dari tim developer agar memenuhi ekspektasi tampilan koordinat & peta lokasi absensi.')
    r3_t.font.name = 'Arial'
    r3_t.font.size = Pt(9.5)
    
    # Row 2, Cell 1 -> Status text
    c2 = table.rows[2].cells[1]
    c2.text = ''
    p_status = c2.paragraphs[0]
    r_lbl = p_status.add_run('Hasil yang diharapkan [STATUS]:\n')
    r_lbl.font.name = 'Arial'
    r_lbl.font.size = Pt(10)
    r_lbl.font.bold = True
    
    expected_desc = 'Berhasil menampilkan informasi Clock In yang mencakup peta lokasi, tanggal, waktu Clock In, deskripsi, lokasi, dan akurasi.' if 'Clock In' in tc_type else 'Berhasil menampilkan detail log absensi Clock Out.'
    
    r_desc = p_status.add_run(expected_desc + ' - ')
    r_desc.font.name = 'Arial'
    r_desc.font.size = Pt(10)
    
    r_st = p_status.add_run('[Pending / Blocker: Render Map Blank]')
    r_st.font.name = 'Arial'
    r_st.font.size = Pt(10)
    r_st.font.bold = True
    r_st.font.color.rgb = RGBColor(185, 28, 28)

set_blocker_note(t86, 'Clock In')
set_blocker_note(t87, 'Clock Out')

doc.save(doc_path)
print('Successfully updated Dokumen_Hasil_Uji_Mobile.docx with clean blocker notes!')
