import openpyxl
from collections import OrderedDict
import copy

EXCEL_LOCAL = r'd:\Project\BTN\SIT\Test Case.xlsx'
EXCEL_DRIVE = r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\Test Case.xlsx'

wb = openpyxl.load_workbook(EXCEL_LOCAL)
ws = wb['TC BTN SMART Web']

# Read all data rows
grouped_rows = OrderedDict()
for r in range(2, ws.max_row + 1):
    m = ws.cell(r, 2).value
    s = ws.cell(r, 3).value
    t = ws.cell(r, 4).value
    if m and str(m).strip() and t and str(t).strip():
        m_str = str(m).strip()
        s_str = str(s).strip() if s else ''
        key = (m_str, s_str)
        if key not in grouped_rows:
            grouped_rows[key] = []
        
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        grouped_rows[key].append(row_vals)

print(f"Total unique submenus: {len(grouped_rows)}")
total_tcs = sum(len(v) for v in grouped_rows.values())
print(f"Total TCs: {total_tcs}")

# Verify 100% count match
assert len(grouped_rows) == 75, f"Expected 75 submenus, got {len(grouped_rows)}"
assert total_tcs == 921, f"Expected 921 TCs, got {total_tcs}"

# Clear existing rows 2..max_row
for r in range(2, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(r, c).value = None

# Write grouped rows with new sec_idx.tc_idx numbering
current_row = 2
for sec_idx, (key, rows_list) in enumerate(grouped_rows.items(), 1):
    mod_name, sub_name = key
    for tc_idx, row_vals in enumerate(rows_list, 1):
        new_no = f"{sec_idx}.{tc_idx}"
        row_vals[0] = new_no # Column 1 (No)
        row_vals[1] = mod_name # Column 2 (Module)
        row_vals[2] = sub_name if sub_name else None # Column 3 (Sub Menu)
        
        for col_idx, val in enumerate(row_vals, 1):
            ws.cell(current_row, col_idx).value = val
        current_row += 1

print(f"Wrote rows up to: {current_row - 1}")

wb.save(EXCEL_LOCAL)
print(f"Saved local: {EXCEL_LOCAL}")
wb.save(EXCEL_DRIVE)
print(f"Saved drive: {EXCEL_DRIVE}")
