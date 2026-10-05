import os
import shutil
import re
from PIL import Image
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def win_p(p):
    ap = os.path.abspath(p)
    return '\\\\?\\' + ap if not ap.startswith('\\\\?\\') else p

d_base = r'D:\Project\BTN Smart\Refactor'
h_base = r'H:\My Drive\Zegen\BTN Smart\Refactor'
g_base = r'G:\My Drive\Zegen\BTN Smart\Refactor'

rel_mod06 = r'Hasil Uji\Screenshot\Web\06. User Authority - User'
rel_mod06_orig = r'Hasil Uji\Screenshot\Web\06. User Authority - User (Original Full)'

src_06_d = os.path.join(d_base, rel_mod06)
orig_06_d = os.path.join(d_base, rel_mod06_orig)

print("=== 1. CREATING ORIGINAL FULL BACKUP FOR MODUL 06 ===")
for base_dir, label in [(d_base, 'Local D'), (h_base, 'Drive H'), (g_base, 'Drive G')]:
    t_orig = os.path.join(base_dir, rel_mod06_orig)
    if not os.path.exists(win_p(t_orig)):
        shutil.copytree(win_p(src_06_d), win_p(t_orig))
        print(f"[{label}] Created Original Full Backup: {t_orig}")
    else:
        print(f"[{label}] Backup already exists!")

print("\n=== 2. CROPPING SCREENSHOTS FOR MODUL 06 ===")
# Crop rules for Modul 06
# Note: 6.1 - 6.8 have sidebar width 300px! 6.9 has sidebar width 338px!
crop_rules_06 = {
    # 6.1 Opening
    ('6.1 Membuka halaman User', '1.png'): (0, 0), # Full screen
    # 6.2 Search valid
    ('6.2 Mencari data User dengan keyword valid', '1.png'): (300, 72),
    # 6.3 Search invalid
    ('6.3 Mencari data User dengan keyword tidak valid', '1.png'): (300, 72),
    # 6.4 Add User
    ('6.4 Menambahkan data User', '1.png'): (300, 72),
    ('6.4 Menambahkan data User', '2.png'): (300, 72),
    ('6.4 Menambahkan data User', '3.png'): (300, 0), # Top alert
    # 6.5 Add Mandatory
    ('6.5 Menambahkan data User tanpa mengisi field mandatory', '1.png'): (300, 72),
    ('6.5 Menambahkan data User tanpa mengisi field mandatory', '2.png'): (300, 72),
    # 6.6 Edit User
    ('6.6 Mengubah data User', '1.png'): (300, 72),
    ('6.6 Mengubah data User', '2.png'): (300, 72),
    ('6.6 Mengubah data User', '3.png'): (300, 0), # Top alert
    # 6.7 Edit Mandatory
    ('6.7 Mengubah data User tanpa mengisi field mandatory', '1.png'): (300, 72),
    ('6.7 Mengubah data User tanpa mengisi field mandatory', '2.png'): (300, 72),
    # 6.8 Delete User
    ('6.8 Menghapus data User', '1.png'): (300, 72),
    ('6.8 Menghapus data User', '2.png'): (300, 0), # Top alert
    # 6.9 Export User
    ('6.9 Mengekspor data User', '1.png'): (338, 72), # Table export
    ('6.9 Mengekspor data User', '2.png'): (338, 72), # Export history
    ('6.9 Mengekspor data User', '3.png'): (0, 0),   # Excel full screen
}

crop_count = 0
for (tc_dir, img_file), (crop_x, crop_y) in crop_rules_06.items():
    src_f = os.path.join(orig_06_d, tc_dir, img_file)
    dst_f = os.path.join(src_06_d, tc_dir, img_file)
    
    with Image.open(win_p(src_f)) as im:
        w, h = im.size
        if crop_x == 0 and crop_y == 0:
            cropped = im.copy()
            action = "FULL SCREEN"
        else:
            cropped = im.crop((crop_x, crop_y, w, h))
            action = f"CROP ({crop_x}, {crop_y})"
        cropped.save(win_p(dst_f))
        crop_count += 1
        print(f"[{crop_count:2d}] {tc_dir[:25]} / {img_file} -> {action} | {cropped.size}")

print(f"\nSuccessfully cropped {crop_count} images for Modul 06!")

print("\n=== 3. EMBEDDING MODUL 06 INTO Dokumen_Hasil_Uji_Web.docx ===")
doc_p = os.path.join(d_base, r'Hasil Uji\Dokumen_Hasil_Uji_Web.docx')
doc = docx.Document(doc_p)

# Find Modul 06 tables
tbl_0 = None
tbl_1 = None
current_h = ''

for el in doc.element.body:
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(el, doc)
        if p.style.name.startswith('Heading 1'):
            current_h = p.text
    elif tag == 'tbl':
        if '6. Modul User Authority - User' in current_h:
            if tbl_0 is None:
                tbl_0 = docx.table.Table(el, doc)
            elif tbl_1 is None:
                tbl_1 = docx.table.Table(el, doc)
                break

print(f"Found Modul 06 tables: Table 0={len(tbl_0.rows)} rows, Table 1={len(tbl_1.rows)} rows")

def natural_sort_key(s):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

def get_images_for_tc(tc_str):
    for d in os.listdir(win_p(src_06_d)):
        if d.startswith(tc_str + ' '):
            fld = os.path.join(src_06_d, d)
            if os.path.isdir(win_p(fld)):
                imgs = [f for f in os.listdir(win_p(fld)) if f.lower().endswith(('.png', '.jpg', '.jpeg'))]
                imgs.sort(key=natural_sort_key)
                return [os.path.join(fld, img) for img in imgs]
    return []

def set_clean_expected_result(cell, desc_text):
    clean_desc = desc_text.strip()
    clean_desc = re.sub(r'^\d+\.\s*', '', clean_desc)
    clean_desc = re.sub(r'\s*-\s*\[Success\]', '', clean_desc).strip()
    
    cell.text = ''
    
    # Paragraph 0: Header
    p0 = cell.paragraphs[0]
    p0_pr = p0._p.get_or_add_pPr()
    p0_sp = OxmlElement('w:spacing')
    p0_sp.set(qn('w:before'), '116')
    p0_sp.set(qn('w:after'), '0')
    p0_pr.append(p0_sp)
    
    r0 = p0.add_run('Hasil yang diharapkan [STATUS]:')
    r0.font.name = 'Arial'
    r0.font.size = Pt(10)
    r0.bold = True
    
    # Paragraph 1: Description + [Success]
    p1 = cell.add_paragraph()
    p1_pr = p1._p.get_or_add_pPr()
    p1_sp = OxmlElement('w:spacing')
    p1_sp.set(qn('w:before'), '60')
    p1_sp.set(qn('w:after'), '100')
    p1_pr.append(p1_sp)
    
    r1 = p1.add_run(clean_desc)
    r1.font.name = 'Arial'
    r1.font.size = Pt(10)
    r1_pr = r1._element.get_or_add_rPr()
    c1 = OxmlElement('w:color')
    c1.set(qn('w:val'), '2C5293')
    r1_pr.append(c1)
    
    r2 = p1.add_run(' - [Success]')
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.bold = True
    r2_pr = r2._element.get_or_add_rPr()
    c2 = OxmlElement('w:color')
    c2.set(qn('w:val'), '2C5293')
    r2_pr.append(c2)

exp_results_06 = {
    '6.1': 'Berhasil menampilkan halaman User.',
    '6.2': 'Berhasil menampilkan data sesuai keyword.',
    '6.3': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '6.4': 'Berhasil menyimpan data user.',
    '6.5': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '6.6': 'Berhasil menyimpan perubahan data user.',
    '6.7': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '6.8': 'Berhasil menghapus data user.',
    '6.9': 'Berhasil mendownload file export data user.',
}

# Embed 6.1 in tbl_0 (row 2 image, row 3 expected)
imgs_6_1 = get_images_for_tc('6.1')
c6_0 = tbl_0.rows[2].cells[1]
c6_0.text = ''
p6_0 = c6_0.paragraphs[0]
p6_0.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, img in enumerate(imgs_6_1):
    p = p6_0 if i == 0 else c6_0.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img, width=Inches(5.9))
set_clean_expected_result(tbl_0.rows[3].cells[1], exp_results_06['6.1'])
print(f"Embedded 6.1: {len(imgs_6_1)} images")

# Embed 6.2 to 6.9 in tbl_1
for k in range(2, 10):
    tc_str = f"6.{k}"
    idx = k - 2
    r_img = idx * 3 + 1
    r_exp = idx * 3 + 2
    
    c = tbl_1.rows[r_img].cells[1]
    c.text = ''
    p_img = c.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    imgs = get_images_for_tc(tc_str)
    for i, img in enumerate(imgs):
        p = p_img if i == 0 else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(5.9))
        
    set_clean_expected_result(tbl_1.rows[r_exp].cells[1], exp_results_06[tc_str])
    print(f"Embedded {tc_str}: {len(imgs)} images | Expected: {exp_results_06[tc_str]}")

doc.save(doc_p)
print(f"\nSuccessfully saved {doc_p} with Modul 06 cropped and embedded!")
