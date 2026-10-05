import os
import shutil
import re
from PIL import Image
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def win_p(p):
    ap = os.path.abspath(p)
    return '\\\\?\\' + ap if not ap.startswith('\\\\?\\') else p

d_base = r'D:\Project\BTN Smart\Refactor'
h_base = r'H:\My Drive\Zegen\BTN Smart\Refactor'
g_base = r'G:\My Drive\Zegen\BTN Smart\Refactor'

rel_mod07 = r'Hasil Uji\Screenshot\Web\07. User Authority - Group Role'
rel_mod07_orig = r'Hasil Uji\Screenshot\Web\07. User Authority - Group Role (Original Full)'

src_07_d = os.path.join(d_base, rel_mod07)
orig_07_d = os.path.join(d_base, rel_mod07_orig)

print("=== 1. CREATING ORIGINAL FULL BACKUP FOR MODUL 07 ===")
for base_dir, label in [(d_base, 'Local D'), (h_base, 'Drive H'), (g_base, 'Drive G')]:
    t_orig = os.path.join(base_dir, rel_mod07_orig)
    if not os.path.exists(win_p(t_orig)):
        shutil.copytree(win_p(src_07_d), win_p(t_orig))
        print(f"[{label}] Created Original Full Backup: {t_orig}")
    else:
        print(f"[{label}] Backup already exists!")

print("\n=== 2. CROPPING SCREENSHOTS FOR MODUL 07 ===")
# Sidebar width is 338px for all normal pages!
crop_rules_07 = {
    # 7.1 Opening
    ('7.1 Membuka halaman Group Role', '1.png'): (0, 0), # Full screen
    # 7.2 Search valid
    ('7.2 Mencari data Group Role dengan keyword valid', '1.png'): (338, 72),
    # 7.3 Search invalid
    ('7.3 Mencari data Group Role dengan keyword tidak valid', '1.png'): (338, 72),
    # 7.4 Add Group Role
    ('7.4 Menambahkan data Group Role', '1.png'): (338, 72),
    ('7.4 Menambahkan data Group Role', '2.png'): (338, 72), # Long scrolled form
    ('7.4 Menambahkan data Group Role', '3.png'): (338, 0),  # Top alert
    # 7.5 Add Mandatory
    ('7.5 Menambahkan data Group Role tanpa mengisi field mandatory', '1.png'): (338, 72),
    ('7.5 Menambahkan data Group Role tanpa mengisi field mandatory', '2.png'): (338, 72), # Long form validation
    # 7.6 Edit Group Role
    ('7.6 Mengubah data Group Role', '1.png'): (338, 72),
    ('7.6 Mengubah data Group Role', '2.png'): (338, 72), # Long form
    ('7.6 Mengubah data Group Role', '3.png'): (338, 0),  # Top alert
    # 7.7 Edit Mandatory
    ('7.7 Mengubah data Group Role tanpa mengisi field mandatory', '1.png'): (338, 72),
    # 7.8 Delete Group Role
    ('7.8 Menghapus data Group Role', '1.png'): (338, 72), # Modal confirmation
    ('7.8 Menghapus data Group Role', '2.png'): (338, 0),  # Top alert
    # 7.9 Export Group Role
    ('7.9 Mengekspor data Group Role', '1.png'): (338, 72), # Table export
    ('7.9 Mengekspor data Group Role', '2.png'): (338, 72), # Export history
    ('7.9 Mengekspor data Group Role', '3.png'): (0, 0),   # Excel full screen
}

crop_count = 0
for (tc_dir, img_file), (crop_x, crop_y) in crop_rules_07.items():
    src_f = os.path.join(orig_07_d, tc_dir, img_file)
    dst_f = os.path.join(src_07_d, tc_dir, img_file)
    
    with Image.open(win_p(src_f)) as im:
        w, h = im.size
        if crop_x == 0 and crop_y == 0:
            cropped = im.copy()
            action = "FULL SCREEN"
        else:
            cropped = im.crop((crop_x, crop_y, w, h))
            action = f"CROP ({crop_x}, {crop_y})"
        cropped.save(win_p(dst_f))
        crop_count += 1
        print(f"[{crop_count:2d}] {tc_dir[:25]} / {img_file} -> {action} | {cropped.size}")

print(f"\nSuccessfully cropped {crop_count} images for Modul 07!")

print("\n=== 3. EMBEDDING MODUL 07 INTO Dokumen_Hasil_Uji_Web.docx ===")
doc_p = os.path.join(d_base, r'Hasil Uji\Dokumen_Hasil_Uji_Web.docx')
doc = docx.Document(doc_p)

# Find Modul 07 tables
tbl_0 = None
tbl_1 = None
current_h = ''

for el in doc.element.body:
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(el, doc)
        if p.style.name.startswith('Heading 1'):
            current_h = p.text
    elif tag == 'tbl':
        if '7. Modul User Authority - Group Role' in current_h:
            if tbl_0 is None:
                tbl_0 = docx.table.Table(el, doc)
            elif tbl_1 is None:
                tbl_1 = docx.table.Table(el, doc)
                break

print(f"Found Modul 07 tables: Table 0={len(tbl_0.rows)} rows, Table 1={len(tbl_1.rows)} rows")

def natural_sort_key(s):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

def get_images_for_tc(tc_str):
    for d in os.listdir(win_p(src_07_d)):
        if d.startswith(tc_str + ' '):
            fld = os.path.join(src_07_d, d)
            if os.path.isdir(win_p(fld)):
                imgs = [f for f in os.listdir(win_p(fld)) if f.lower().endswith(('.png', '.jpg', '.jpeg'))]
                imgs.sort(key=natural_sort_key)
                return [os.path.join(fld, img) for img in imgs]
    return []

def set_clean_expected_result(cell, desc_text):
    clean_desc = desc_text.strip()
    clean_desc = re.sub(r'^\d+\.\s*', '', clean_desc)
    clean_desc = re.sub(r'\s*-\s*\[Success\]', '', clean_desc).strip()
    
    cell.text = ''
    
    # Paragraph 0: Header
    p0 = cell.paragraphs[0]
    p0_pr = p0._p.get_or_add_pPr()
    p0_sp = OxmlElement('w:spacing')
    p0_sp.set(qn('w:before'), '116')
    p0_sp.set(qn('w:after'), '0')
    p0_pr.append(p0_sp)
    
    r0 = p0.add_run('Hasil yang diharapkan [STATUS]:')
    r0.font.name = 'Arial'
    r0.font.size = Pt(10)
    r0.bold = True
    
    # Paragraph 1: Description + [Success]
    p1 = cell.add_paragraph()
    p1_pr = p1._p.get_or_add_pPr()
    p1_sp = OxmlElement('w:spacing')
    p1_sp.set(qn('w:before'), '60')
    p1_sp.set(qn('w:after'), '100')
    p1_pr.append(p1_sp)
    
    r1 = p1.add_run(clean_desc)
    r1.font.name = 'Arial'
    r1.font.size = Pt(10)
    r1_pr = r1._element.get_or_add_rPr()
    c1 = OxmlElement('w:color')
    c1.set(qn('w:val'), '2C5293')
    r1_pr.append(c1)
    
    r2 = p1.add_run(' - [Success]')
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.bold = True
    r2_pr = r2._element.get_or_add_rPr()
    c2 = OxmlElement('w:color')
    c2.set(qn('w:val'), '2C5293')
    r2_pr.append(c2)

exp_results_07 = {
    '7.1': 'Berhasil menampilkan halaman Group Role.',
    '7.2': 'Berhasil menampilkan data sesuai keyword.',
    '7.3': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '7.4': 'Berhasil menyimpan data group role.',
    '7.5': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '7.6': 'Berhasil menyimpan perubahan data group role.',
    '7.7': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '7.8': 'Berhasil menghapus data group role.',
    '7.9': 'Berhasil mendownload file export data group role.',
}

# Embed 7.1 in tbl_0 (row 2 image, row 3 expected)
imgs_7_1 = get_images_for_tc('7.1')
c7_0 = tbl_0.rows[2].cells[1]
c7_0.text = ''
p7_0 = c7_0.paragraphs[0]
p7_0.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, img in enumerate(imgs_7_1):
    p = p7_0 if i == 0 else c7_0.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img, width=Inches(5.9))
set_clean_expected_result(tbl_0.rows[3].cells[1], exp_results_07['7.1'])
print(f"Embedded 7.1: {len(imgs_7_1)} images")

# Embed 7.2 to 7.9 in tbl_1
for k in range(2, 10):
    tc_str = f"7.{k}"
    idx = k - 2
    r_img = idx * 3 + 1
    r_exp = idx * 3 + 2
    
    c = tbl_1.rows[r_img].cells[1]
    c.text = ''
    p_img = c.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    imgs = get_images_for_tc(tc_str)
    for i, img in enumerate(imgs):
        p = p_img if i == 0 else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(5.9))
        
    set_clean_expected_result(tbl_1.rows[r_exp].cells[1], exp_results_07[tc_str])
    print(f"Embedded {tc_str}: {len(imgs)} images | Expected: {exp_results_07[tc_str]}")

doc.save(doc_p)
print(f"\nSuccessfully saved {doc_p} with Modul 07 cropped and embedded!")
