import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, re, shutil

doc_path = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
img_base = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile'

doc = docx.Document(doc_path)

tcs = [
    ('7.1', True),
    ('7.2', False),
    ('11.1', True),
    ('11.2', False),
    ('11.3', False)
]

for tc_num, is_first in tcs:
    print(f"\n=== Processing TC {tc_num} ===")
    target_table = None
    for tbl in doc.tables:
        r0 = tbl.rows[0].cells[0].text.strip()
        if not is_first and r0 == tc_num:
            target_table = tbl
            break
        elif is_first and r0 == 'No.' and len(tbl.rows) > 1 and tbl.rows[1].cells[0].text.strip() == tc_num:
            target_table = tbl
            break
            
    if not target_table:
        print(f"Error: Table for {tc_num} not found!")
        continue
        
    image_cell = target_table.rows[2].cells[1] if is_first else target_table.rows[1].cells[1]
    exp_cell = target_table.rows[3].cells[1] if is_first else target_table.rows[2].cells[1]
    
    tc_folder = None
    for root, dirs, files in os.walk(img_base):
        for d in dirs:
            if d.startswith(tc_num + ' ') or d == tc_num:
                tc_folder = os.path.join(root, d)
                break
        if tc_folder: break
        
    if not tc_folder:
        print(f"Error: Folder for {tc_num} not found!")
        continue
        
    raw_images = [f for f in os.listdir(tc_folder) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
    numbered = [f for f in raw_images if re.match(r'^\d+\.', f)]
    if numbered:
        sorted_images = sorted(numbered, key=lambda x: int(re.match(r'^(\d+)', x).group(1)))
    else:
        sorted_images = sorted(raw_images)
        
    full_paths = [os.path.join(tc_folder, f) for f in sorted_images]
    print(f"  Folder: {os.path.basename(tc_folder)}")
    print(f"  Images: {sorted_images}")
    
    # Clear existing paragraphs in image cell
    for p in list(image_cell.paragraphs):
        p._element.getparent().remove(p._element)
        
    p_img = image_cell.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(6)
    run_img = p_img.add_run()
    
    if len(full_paths) == 1:
        h = 4.0
    elif len(full_paths) == 2:
        h = 3.8
    elif len(full_paths) == 3:
        h = 3.2
    else:
        h = 2.8
        
    for idx, imp in enumerate(full_paths):
        run_img.add_picture(imp, height=Inches(h))
        if idx < len(full_paths) - 1:
            run_img.add_text('   ')
    print(f"  Inserted {len(full_paths)} images side-by-side (height={h} in)")
    
    # Status
    if '[Success]' not in exp_cell.text:
        last_p = exp_cell.paragraphs[-1]
        run_status = last_p.add_run(" - [Success]")
        run_status.font.name = 'Arial'
        run_status.font.size = Pt(10)
        run_status.bold = True
        run_status.font.color.rgb = RGBColor(0x2C, 0x52, 0x93)
        print("  Appended - [Success]")
    else:
        print("  [Success] already present")

doc.save(doc_path)
print(f"\nSaved successfully to {doc_path}!")

# Sync to Drive H and Drive G
for dst in [
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
]:
    if os.path.exists(os.path.dirname(dst)):
        shutil.copy2(doc_path, dst)
        print(f"Synced to: {dst}")
