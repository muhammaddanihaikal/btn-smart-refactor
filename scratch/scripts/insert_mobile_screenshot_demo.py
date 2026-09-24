import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, shutil

DOC_PATH = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
IMG_PATH = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile\01. Login\1.1 Membuka halaman Login\2.jpg'

print(f"Loading document: {DOC_PATH}")
doc = docx.Document(DOC_PATH)

target_table = None
for i, tbl in enumerate(doc.tables):
    for r in tbl.rows:
        if len(r.cells) >= 2 and r.cells[0].text.strip() == '1.1':
            target_table = tbl
            print(f"Found TC 1.1 at Table index {i}")
            break
    if target_table:
        break

if not target_table:
    print("Error: Table for TC 1.1 not found!")
    exit(1)

# Row 2 Cell 1 is the image cell
img_cell = target_table.rows[2].cells[1]
# Clear existing text/paragraphs
p_elem = img_cell.paragraphs[0]
for p in img_cell.paragraphs:
    p_elem_to_remove = p._element
    p_elem_to_remove.getparent().remove(p_elem_to_remove)

# Add centered paragraph for image
p_img = img_cell.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.space_before = Pt(6)
p_img.paragraph_format.space_after = Pt(6)

run_img = p_img.add_run()
# Portrait mobile screenshot: 1220x2530 aspect ratio ~ 0.482
# Setting height = 4.2 inches gives width ~ 2.03 inches, perfectly framed!
run_img.add_picture(IMG_PATH, height=Inches(4.2))
print("Inserted 2.jpg into Row 2 Cell 1.")

# Row 3 Cell 1 is the Expected Results & Status
exp_cell = target_table.rows[3].cells[1]
# Check paragraphs in exp_cell
if len(exp_cell.paragraphs) > 1:
    p_status = exp_cell.paragraphs[1]
    if '[Success]' not in p_status.text:
        # Check text
        txt = p_status.text.strip()
        # clear runs
        for r in list(p_status.runs):
            p_status._element.remove(r._element)
        run_txt = p_status.add_run(txt + " - [Success]")
        run_txt.font.color.rgb = RGBColor(0x2C, 0x52, 0x93)
        print("Updated status to [Success] with blue font.")
elif len(exp_cell.paragraphs) == 1:
    txt = exp_cell.paragraphs[0].text.strip()
    if '[Success]' not in txt:
        run_txt = exp_cell.paragraphs[0].add_run(" - [Success]")
        run_txt.font.color.rgb = RGBColor(0x2C, 0x52, 0x93)
        print("Appended [Success] to status.")

# Save document
doc.save(DOC_PATH)
print(f"Saved changes to: {DOC_PATH}")

# Sync to local and Drive G
for dst in [
    r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
]:
    if os.path.exists(os.path.dirname(dst)):
        shutil.copy2(DOC_PATH, dst)
        print(f"Synced to: {dst}")

print("ALL DONE!")
