import openpyxl, shutil

xlsx_path = r'D:\Project\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx'
wb = openpyxl.load_workbook(xlsx_path)
ws = wb['TC BTN SMART Mobile']

# Verification before deletion
print("Pre-check:")
print("Row 150:", ws.cell(150, 1).value, ws.cell(150, 4).value)
print("Row 144:", ws.cell(144, 1).value, ws.cell(144, 4).value)
print("Row 80 :", ws.cell(80, 1).value, ws.cell(80, 4).value)

# 1. Delete Row 150 (Daily Sales Agenda - Masuk)
assert ws.cell(150, 1).value == '15.2' and 'masuk' in str(ws.cell(150, 4).value).lower()
ws.delete_rows(150)
print("\nDeleted row 150: 15.2 Melihat daftar agenda pada tab Masuk")

# Renumber remaining Modul 15 rows (now rows 150 to 155)
for i in range(150, 156):
    new_no = f"15.{i - 148}"
    ws.cell(i, 1).value = new_no
    print(f"  Row {i}: set No to {new_no} ({ws.cell(i, 4).value})")

# 2. Delete Row 144 (Pengingat - Aktif)
assert ws.cell(144, 1).value == '14.2' and 'aktif' in str(ws.cell(144, 4).value).lower()
ws.delete_rows(144)
print("\nDeleted row 144: 14.2 Melihat daftar pengingat pada tab Aktif")

# Renumber remaining Modul 14 rows (now rows 144 to 147)
for i in range(144, 148):
    new_no = f"14.{i - 142}"
    ws.cell(i, 1).value = new_no
    print(f"  Row {i}: set No to {new_no} ({ws.cell(i, 4).value})")

# 3. Delete Row 80 (Cuti & Izin - Menampilkan daftar)
assert ws.cell(80, 1).value == '9.2' and 'menampilkan daftar' in str(ws.cell(80, 4).value).lower()
ws.delete_rows(80)
print("\nDeleted row 80: 9.2 Menampilkan daftar data pengajuan cuti")

# Renumber remaining Modul 09 rows (now rows 80 to 86)
for i in range(80, 87):
    new_no = f"9.{i - 78}"
    ws.cell(i, 1).value = new_no
    print(f"  Row {i}: set No to {new_no} ({ws.cell(i, 4).value})")

wb.save(xlsx_path)
print(f"\nSaved successfully to: {xlsx_path}")

# Sync to Drive H and Drive G
for dst in [
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx'
]:
    shutil.copy2(xlsx_path, dst)
    print(f"Synced Test Script BTN Smart Refactor.xlsx to: {dst}")
