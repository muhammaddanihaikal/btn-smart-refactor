import openpyxl
import docx
from collections import Counter

# 1. Read Excel
wb = openpyxl.load_workbook(r'D:\Project\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx', data_only=True)
ws = wb['TC BTN SMART Mobile']

print("=== EXCEL VISIT & ATTENDANCE ROWS ===")
for r in range(2, ws.max_row + 1):
    mod = ws.cell(r, 2).value
    sub = ws.cell(r, 3).value
    title = ws.cell(r, 4).value
    scen = ws.cell(r, 5).value
    exp = ws.cell(r, 6).value
    if title and any(k in str(title).lower() for k in ['visit', 'clock', 'absen']):
        print(f"Row {r:3d} | Mod: {str(mod):22s} | Sub: {str(sub):15s} | Title: {title}")

# 2. Read Senior Word doc
doc_senior = docx.Document(r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile (Updated).docx')
titles_senior = []
for tbl in doc_senior.tables:
    if len(tbl.columns) == 7:
        for r in tbl.rows:
            c0 = r.cells[0].text.strip()
            if c0 and '.' in c0 and c0[0].isdigit():
                t = r.cells[1].text.strip()
                sc = r.cells[2].text.strip()
                ex = r.cells[3].text.strip()
                titles_senior.append((c0, t, sc, ex))

print("\n=== DUPLICATE TEST CASE TITLES IN SENIOR DOC ===")
c = Counter([t[1] for t in titles_senior])
for title, cnt in c.items():
    if cnt > 1:
        items = [t for t in titles_senior if t[1] == title]
        print(f"\nTitle: '{title}' ({cnt}x)")
        for item in items:
            print(f"   No {item[0]:5s} | Scenario: {item[2][:60]}... | Expected: {item[3][:60]}...")

# 3. Read Generated SIT Doc
doc_gen = docx.Document(r'D:\Project\BTN Smart\Refactor\SIT\Document SIT BTN Smart\SIT BTN SMART Mobile - 24 Sep 2026.docx')
current_mod = ""
print("\n=== CURRENT GENERATED MOBILE SIT DOC: VISIT & ATTENDANCE ===")
for child in doc_gen.element.body:
    if child.tag.endswith('p'):
        p = docx.text.paragraph.Paragraph(child, doc_gen)
        txt = p.text.strip()
        if txt.startswith('Modul') or ('Modul ' in txt and any(c.isdigit() for c in txt[:5])):
            current_mod = txt
    elif child.tag.endswith('tbl'):
        tbl = docx.table.Table(child, doc_gen)
        if len(tbl.columns) == 7:
            for r in tbl.rows:
                c0 = r.cells[0].text.strip()
                if c0 and '.' in c0 and c0[0].isdigit():
                    t = r.cells[1].text.strip()
                    if any(k in t.lower() or k in current_mod.lower() for k in ['visit', 'clock', 'absen']):
                        print(f"No {c0:6s} | {current_mod:45s} | Title: {t}")
