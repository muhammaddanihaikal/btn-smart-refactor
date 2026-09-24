import docx

doc_path = r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx'
doc = docx.Document(doc_path)
print(f"Total tables in SIT Web: {len(doc.tables)}")

for tbl in doc.tables[:3]:
    for r in tbl.rows[:5]:
        print([c.text.strip().replace('\n', ' ') for c in r.cells])
