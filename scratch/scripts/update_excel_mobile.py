import openpyxl, shutil

xlsx_path = r'D:\Project\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx'
wb = openpyxl.load_workbook(xlsx_path)
ws = wb['TC BTN SMART Mobile']

# Verify the rows before deleting
print("Row 150:", ws.cell(150, 3).value)
print("Row 144:", ws.cell(144, 3).value)
print("Row 80 :", ws.cell(80, 3).value)

# Delete in reverse order
assert 'masuk' in str(ws.cell(150, 3).value).lower()
ws.delete_rows(150)
print("Deleted row 150 (Melihat daftar agenda pada tab Masuk)")

assert 'aktif' in str(ws.cell(144, 3).value).lower()
ws.delete_rows(144)
print("Deleted row 144 (Melihat daftar pengingat pada tab Aktif)")

assert 'cuti' in str(ws.cell(80, 3).value).lower()
ws.delete_rows(80)
print("Deleted row 80 (Menampilkan daftar data pengajuan cuti)")

wb.save(xlsx_path)
print(f"Saved: {xlsx_path}")

# Sync to Drive H and Drive G
for dst in [
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx'
]:
    shutil.copy2(xlsx_path, dst)
    print(f"Synced Excel to: {dst}")
