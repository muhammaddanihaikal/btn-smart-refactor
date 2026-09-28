import docx, shutil

doc_path = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
doc = docx.Document(doc_path)

print(f"Total tables before: {len(doc.tables)}")

# Identify tables to delete by exact TC number and Title
del_targets = [
    ('9.2', 'Menampilkan daftar data pengajuan cuti'),
    ('14.2', 'Melihat daftar pengingat pada tab Aktif'),
    ('15.2', 'Melihat daftar agenda pada tab Masuk')
]

deleted_count = 0
for tbl in list(doc.tables):
    r0_c0 = tbl.rows[0].cells[0].text.strip()
    r0_c1 = tbl.rows[0].cells[1].text.strip().replace('\n', ' ') if len(tbl.rows[0].cells) > 1 else ''
    
    # Check if this table matches any delete target
    for tc_no, tc_title in del_targets:
        if r0_c0 == tc_no and tc_title.lower() in r0_c1.lower():
            print(f"Deleting table: {tc_no} - {r0_c1}")
            tbl._element.getparent().remove(tbl._element)
            deleted_count += 1
            break

print(f"Deleted {deleted_count} tables. Remaining tables: {len(doc.tables)}")

# Now renumber the subsequent tables in Modul 09, 14, 15
renumber_map = {
    # Modul 09
    '9.3': '9.2',
    '9.4': '9.3',
    '9.5': '9.4',
    '9.6': '9.5',
    '9.7': '9.6',
    '9.8': '9.7',
    '9.9': '9.8',
    # Modul 14
    '14.3': '14.2',
    '14.4': '14.3',
    '14.5': '14.4',
    '14.6': '14.5',
    # Modul 15
    '15.3': '15.2',
    '15.4': '15.3',
    '15.5': '15.4',
    '15.6': '15.5',
    '15.7': '15.6',
    '15.8': '15.7',
}

renumbered_count = 0
for tbl in doc.tables:
    r0_c0 = tbl.rows[0].cells[0].text.strip()
    if r0_c0 in renumber_map:
        new_no = renumber_map[r0_c0]
        tbl.rows[0].cells[0].text = new_no
        title = tbl.rows[0].cells[1].text.strip().replace('\n', ' ') if len(tbl.rows[0].cells) > 1 else ''
        print(f"Renumbered: {r0_c0} -> {new_no} ({title})")
        renumbered_count += 1

print(f"Total renumbered: {renumbered_count}")

doc.save(doc_path)
print(f"Saved changes to: {doc_path}")

# Verify doc can be opened cleanly
test_doc = docx.Document(doc_path)
print(f"Verification: Document opened cleanly! Total tables: {len(test_doc.tables)}")

# Sync to Drive H and Drive G
for dst in [
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
]:
    shutil.copy2(doc_path, dst)
    print(f"Synced Hasil Uji to: {dst}")
