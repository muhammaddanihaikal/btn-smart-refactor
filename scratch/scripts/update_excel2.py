import openpyxl
import shutil
import os

updates = {
    'Melakukan login dengan email dan password valid': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email terdaftar.\n3. Isi field Password dengan password yang valid.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Berhasil melakukan login dan diarahkan ke halaman dashboard."
    },
    'Melakukan login dengan email tidak terdaftar': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email yang tidak terdaftar.\n3. Isi field Password dengan password bebas.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error email tidak terdaftar."
    },
    'Melakukan login dengan password salah': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email terdaftar.\n3. Isi field Password dengan password yang salah.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error password salah."
    },
    'Melakukan login tanpa mengisi email': {
        'steps': "1. Pilih Type.\n2. Kosongkan field Email.\n3. Isi field Password dengan password yang valid.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengosongkan Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error field mandatory."
    },
    'Melakukan login tanpa mengisi password': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email terdaftar.\n3. Kosongkan field Password.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengosongkan Password.\n4. Menampilkan pesan error field mandatory."
    },
    'Melakukan login tanpa mengisi email dan password': {
        'steps': "1. Pilih Type.\n2. Kosongkan field Email.\n3. Kosongkan field Password.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengosongkan Email.\n3. Berhasil mengosongkan Password.\n4. Menampilkan pesan error field mandatory."
    }
}

def update_excel_by_title(file_path):
    print(f"Updating by title {file_path}")
    wb = openpyxl.load_workbook(file_path)
    if 'TC BTN SMART Mobile' not in wb.sheetnames:
        print("  Sheet 'TC BTN SMART Mobile' not found.")
        return
    ws = wb['TC BTN SMART Mobile']
    
    # find column indices
    header_row = 1
    col_title = col_steps = col_expected = -1
    for i, cell in enumerate(ws[header_row]):
        val = str(cell.value).strip().lower()
        if 'title' in val or 'skenario' in val:
            col_title = i + 1
        elif 'langkah-langkah' in val or 'steps' in val:
            col_steps = i + 1
        elif 'hasil yang diharapkan' in val or 'expected' in val:
            col_expected = i + 1
            
    if col_title == -1 or col_steps == -1 or col_expected == -1:
        print(f"  Headers not found! title:{col_title}, steps:{col_steps}, exp:{col_expected}")
        return
        
    updated = 0
    for row in range(2, ws.max_row + 1):
        title_val = str(ws.cell(row, col_title).value).strip()
        if title_val in updates:
            ws.cell(row, col_steps).value = updates[title_val]['steps']
            ws.cell(row, col_expected).value = updates[title_val]['expected']
            updated += 1
            
    if updated > 0:
        wb.save(file_path)
        print(f"  Updated {updated} rows and saved.")
    else:
        print("  No rows updated.")

excel_files2 = [
    r'D:\Project\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx',
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx'
]

for ef in excel_files2:
    if os.path.exists(ef):
        update_excel_by_title(ef)
        # Sync to D and G
        base_name = os.path.basename(ef)
        for drive in ['D', 'G']:
            dst = rf"{drive}:\Project\BTN Smart\Refactor\Test Script\{base_name}" if drive == 'D' else rf"{drive}:\My Drive\Zegen\BTN Smart\Refactor\Test Script\{base_name}"
            if os.path.exists(os.path.dirname(dst)):
                shutil.copy2(ef, dst)
                print(f"  Synced to {dst}")
