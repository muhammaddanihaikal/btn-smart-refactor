import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, glob, re, shutil

DOC_PATH = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
SCREENSHOT_BASE = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile'

def insert_tc_screenshots(tc_list):
    if isinstance(tc_list, str):
        tc_list = [tc_list]
        
    doc = docx.Document(DOC_PATH)
    
    for tc_num in tc_list:
        print(f"\n--- Processing TC {tc_num} ---")
        target_table = None
        is_first_tc = False
        for i, tbl in enumerate(doc.tables):
            r0_c0 = tbl.rows[0].cells[0].text.strip()
            if r0_c0 == tc_num:
                target_table = tbl
                is_first_tc = False
                print(f"Found TC {tc_num} at Table {i}")
                break
            elif r0_c0 == 'No.' and len(tbl.rows) > 1 and tbl.rows[1].cells[0].text.strip() == tc_num:
                target_table = tbl
                is_first_tc = True
                print(f"Found TC {tc_num} (first in module) at Table {i}")
                break
                
        if not target_table:
            print(f"Error: Table for TC {tc_num} not found!")
            continue
            
        if is_first_tc:
            image_cell = target_table.rows[2].cells[1]
            exp_cell = target_table.rows[3].cells[1]
        else:
            image_cell = target_table.rows[1].cells[1]
            exp_cell = target_table.rows[2].cells[1]
            
        tc_folder = None
        for root, dirs, files in os.walk(SCREENSHOT_BASE):
            for d in dirs:
                if d.startswith(tc_num + " ") or d == tc_num:
                    tc_folder = os.path.join(root, d)
                    break
            if tc_folder:
                break
                
        if not tc_folder:
            print(f"Error: Folder for TC {tc_num} not found in {SCREENSHOT_BASE}")
            continue
            
        print(f"Folder: {tc_folder}")
        
        valid_exts = ('.jpg', '.jpeg', '.png')
        raw_images = [f for f in os.listdir(tc_folder) if f.lower().endswith(valid_exts)]
        numbered_images = [f for f in raw_images if re.match(r'^\d+\.', f)]
        if numbered_images:
            images_to_use = sorted(numbered_images, key=lambda x: int(re.match(r'^(\d+)', x).group(1)))
        else:
            images_to_use = sorted(raw_images)
            
        full_image_paths = [os.path.join(tc_folder, f) for f in images_to_use]
        print(f"Images to insert: {images_to_use}")
        
        if not full_image_paths:
            print(f"Warning: No images found in {tc_folder}")
            continue
            
        for p in list(image_cell.paragraphs):
            p._element.getparent().remove(p._element)
            
        p_img = image_cell.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        
        if len(full_image_paths) == 1:
            height_in = 4.0
        elif len(full_image_paths) == 2:
            height_in = 3.8
        elif len(full_image_paths) == 3:
            height_in = 3.2
        else:
            height_in = 2.8
            
        for idx, img_path in enumerate(full_image_paths):
            run_img.add_picture(img_path, height=Inches(height_in))
            if idx < len(full_image_paths) - 1:
                run_img.add_text("   ")
                
        print(f"Inserted {len(full_image_paths)} images side-by-side (height={height_in}\").")
        
        if '[Success]' not in exp_cell.text:
            last_p = exp_cell.paragraphs[-1]
            run_status = last_p.add_run(" - [Success]")
            run_status.font.name = 'Arial'
            run_status.font.size = Pt(10)
            run_status.bold = True
            run_status.font.color.rgb = RGBColor(0x2C, 0x52, 0x93)
            print("Appended [Success] in blue color.")
        else:
            print("[Success] already present.")
            
    doc.save(DOC_PATH)
    print(f"\nSaved changes to: {DOC_PATH}")
    
    for dst in [
        r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
        r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
    ]:
        if os.path.exists(os.path.dirname(dst)):
            shutil.copy2(DOC_PATH, dst)
            print(f"Synced to: {dst}")

tcs_to_process = ['5.1', '5.2', '5.3', '5.4', '5.5', '5.6', '5.7', '6.1', '6.2', '6.3']
insert_tc_screenshots(tcs_to_process)
