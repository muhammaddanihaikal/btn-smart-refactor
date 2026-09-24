import os, shutil, re, copy
import docx
from docx import Document
from docx.table import Table
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import zipfile
import lxml.etree as etree

MOBILE_DOCX   = r'D:\Project\BTN\SIT\Document SIT BTN Smart\SIT BTN SMART Mobile.docx'
TEMPLATE_DOCX = r'D:\Project\BTN\Hasil Uji\Template Dokumen Hasil Uji.docx'
OUTPUT_DOCX   = r'D:\Project\BTN\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
SCREENSHOT_DIR = r'D:\Project\BTN\Hasil Uji\Screenshot\Mobile'

def sanitize_folder_name(name):
    clean = re.sub(r'[<>:"/\\|?*]', '-', str(name))
    clean = re.sub(r'\s+', ' ', clean).strip(' .')
    if len(clean) > 100:
        clean = clean[:100].strip(' .')
    return clean

def clear_cell(cell):
    for p in cell.paragraphs:
        p._element.getparent().remove(p._element)
    cell.add_paragraph()

def set_cell_text(cell, text, bold=False, alignment=None, space_before=None):
    p = cell.paragraphs[0]
    if alignment is not None:
        p.alignment = alignment
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.bold = bold
    return p

# ── 1. PARSE TCs from Mobile SIT ─────────────────────────────────────
doc_src = docx.Document(MOBILE_DOCX)
tcs = []
current_mod = ''
mod_index = 0
last_mod = ''

for child in doc_src.element.body:
    if child.tag.endswith('p'):
        p = docx.text.paragraph.Paragraph(child, doc_src)
        txt = p.text.strip()
        if 'Modul ' in txt and txt[0].isdigit():
            current_mod = txt.split('Modul ')[-1].strip()
            if current_mod != last_mod:
                mod_index += 1
                last_mod = current_mod
    elif child.tag.endswith('tbl'):
        tbl = docx.table.Table(child, doc_src)
        if len(tbl.columns) == 7:
            for r in tbl.rows:
                c0 = r.cells[0].text.strip()
                if c0 and '.' in c0 and c0[0].isdigit():
                    tcs.append({
                        'mod_idx': mod_index,
                        'mod_name': current_mod,
                        'no': c0,
                        'title': r.cells[1].text.strip(),
                        'expected': r.cells[3].text.strip()
                    })

print(f"Extracted {len(tcs)} TCs across {mod_index} submenus.")

# ── 2. LOAD TEMPLATE & IDENTIFY TC TABLES ────────────────────────────
doc_tpl = Document(TEMPLATE_DOCX)

# Find ref_table_first (4 rows: header row + title row + image row + expected row)
# and ref_table_next (3 rows: title row + image row + expected row)
ref_table_first = None  # Table index 4 (has "No." header row)
ref_table_next  = None  # Table index 5 (starts directly with TC number)

for i, tbl in enumerate(doc_tpl.tables):
    rows = tbl.rows
    cols = len(tbl.columns)
    if cols == 2 and len(rows) >= 3:
        c0_0 = rows[0].cells[0].text.strip()
        c0_1 = rows[0].cells[1].text.strip()
        if c0_0 == 'No.' and 'UAT' in c0_1:
            if ref_table_first is None:
                ref_table_first = tbl._tbl
                print(f"ref_table_first found at table {i}")
        elif c0_0 and '.' in c0_0 and c0_0[0].isdigit():
            if ref_table_next is None:
                ref_table_next = tbl._tbl
                print(f"ref_table_next found at table {i}")
    if ref_table_first is not None and ref_table_next is not None:
        break

if ref_table_first is None or ref_table_next is None:
    raise RuntimeError("Could not find TC template tables!")

# ── 3. FIND HEADING REFERENCE AND CLEAR BODY ─────────────────────────
body = doc_tpl.element.body
start_delete_idx = -1
ref_heading_p = None

for i, el in enumerate(body):
    if el.tag == qn('w:p'):
        text = ''.join(t.text or '' for t in el.findall('.//' + qn('w:t')))
        if 'Departemen' in text:
            start_delete_idx = i
            ref_heading_p = copy.deepcopy(el)
            break

if start_delete_idx != -1:
    for el in list(body)[start_delete_idx:]:
        body.remove(el)

# ── 4. BUILD DOCUMENT ─────────────────────────────────────────────────
current_sec = None
for tc in tcs:
    sec_num = tc['mod_idx']

    # Add section heading on new module
    if sec_num != current_sec:
        current_sec = sec_num
        new_p = copy.deepcopy(ref_heading_p)
        for child in list(new_p):
            if child.tag != qn('w:pPr'):
                new_p.remove(child)
        pPr = new_p.find(qn('w:pPr'))
        if pPr is not None:
            numPr = pPr.find(qn('w:numPr'))
            if numPr is not None:
                pPr.remove(numPr)
        r_run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rPr.append(OxmlElement('w:b'))
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '32'); rPr.append(sz)
        szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), '32'); rPr.append(szCs)
        r_run.append(rPr)
        t_el = OxmlElement('w:t')
        t_el.text = f"{sec_num}. Modul {tc['mod_name']}"
        r_run.append(t_el)
        new_p.append(r_run)
        body.append(new_p)
        body.append(OxmlElement('w:p'))
        is_first_tc = True
    else:
        is_first_tc = False

    # Choose template table
    if is_first_tc:
        new_tbl = copy.deepcopy(ref_table_first)
        t = Table(new_tbl, doc_tpl)
        # Row 0: Header "No." | "UAT" -> keep as is
        # Row 1: TC No | TC Title
        clear_cell(t.rows[1].cells[0])
        set_cell_text(t.rows[1].cells[0], tc['no'], space_before=5.8, alignment=1)
        clear_cell(t.rows[1].cells[1])
        set_cell_text(t.rows[1].cells[1], f" {tc['title']}", space_before=5.8)
        image_cell    = t.rows[2].cells[1]
        expected_cell = t.rows[3].cells[1]
    else:
        new_tbl = copy.deepcopy(ref_table_next)
        t = Table(new_tbl, doc_tpl)
        # Row 0: TC No | TC Title
        clear_cell(t.rows[0].cells[0])
        set_cell_text(t.rows[0].cells[0], tc['no'], space_before=5.8, alignment=1)
        clear_cell(t.rows[0].cells[1])
        set_cell_text(t.rows[0].cells[1], f" {tc['title']}", space_before=5.8)
        image_cell    = t.rows[1].cells[1]
        expected_cell = t.rows[2].cells[1]

    # Set Expected Result
    clear_cell(expected_cell)
    p1 = expected_cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(5.8)
    run1 = p1.add_run("Hasil yang diharapkan [STATUS]:\n")
    run1.font.name = 'Arial'; run1.font.size = Pt(10); run1.bold = True
    run2 = p1.add_run(tc['expected'])
    run2.font.name = 'Arial'; run2.font.size = Pt(10)

    # Set Images if exist
    clear_cell(image_cell)
    image_p = image_cell.paragraphs[0]
    image_p.alignment = 1
    image_p.paragraph_format.space_before = Pt(5.8)

    folder_sub = f"{tc['mod_idx']:02d}. {sanitize_folder_name(tc['mod_name'])}"
    folder_tc  = f"{tc['no']} {sanitize_folder_name(tc['title'])}"
    tc_dir = os.path.join(SCREENSHOT_DIR, folder_sub, folder_tc)

    if os.path.exists(tc_dir):
        valid_exts = ('.png', '.jpg', '.jpeg')
        raw_files = [f for f in os.listdir(tc_dir) if f.lower().endswith(valid_exts)]
        files = sorted(raw_files, key=lambda x: [int(c) if c.isdigit() else c for c in re.split(r'(\d+)', x)])
        for f in files:
            img_path = os.path.join(tc_dir, f)
            try:
                run = image_p.add_run()
                run.add_picture(img_path, width=Inches(5.9))
                run.add_break()
            except Exception as e:
                print(f"Error inserting image {img_path}: {e}")

    body.append(new_tbl)
    body.append(OxmlElement('w:p'))

doc_tpl.save(OUTPUT_DOCX)
print(f"Saved: {OUTPUT_DOCX}")

# ── 5. STRIP DUPLICATE PARA IDs ──────────────────────────────────────
tmp = OUTPUT_DOCX + '.tmp'
with zipfile.ZipFile(OUTPUT_DOCX, 'r') as z_in, zipfile.ZipFile(tmp, 'w') as z_out:
    for item in z_in.infolist():
        if item.filename == 'word/document.xml':
            content = z_in.read(item.filename)
            try:
                root = etree.fromstring(content)
                for el in root.iter():
                    for k in list(el.attrib.keys()):
                        if 'paraId' in k or 'textId' in k:
                            del el.attrib[k]
                content = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone='yes')
            except Exception as e:
                print(f"Warning strip IDs: {e}")
            z_out.writestr(item, content)
        else:
            z_out.writestr(item, z_in.read(item.filename))
shutil.move(tmp, OUTPUT_DOCX)
print("Stripped duplicate IDs.")

# ── 6. SYNC TO DRIVE H & G ───────────────────────────────────────────
for drive in ['H', 'G']:
    dst = rf"{drive}:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx"
    try:
        shutil.copy2(OUTPUT_DOCX, dst)
        print(f"Synced to Drive {drive}: {dst}")
    except Exception as e:
        print(f"Failed to sync to Drive {drive}: {e}")

print("\nDone!")
