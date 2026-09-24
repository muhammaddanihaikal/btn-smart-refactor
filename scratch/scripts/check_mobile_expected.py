import docx

doc_path = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
doc = docx.Document(doc_path)
print(f"Total tables in Mobile: {len(doc.tables)}")

sample_expected = []
for i, tbl in enumerate(doc.tables[4:15]):
    for r in tbl.rows:
        for c in r.cells:
            if "Hasil yang diharapkan" in c.text:
                sample_expected.append((i, tbl.rows[0].cells[0].text.strip(), c.text.strip()))
                break

for s in sample_expected:
    print(f"Table {s[0]} ({s[1]}): {repr(s[2])}\n")
