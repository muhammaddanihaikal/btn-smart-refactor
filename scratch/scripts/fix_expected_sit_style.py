import docx
from docx.shared import Pt, RGBColor
import shutil, os

doc_path = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
print("Loading doc...")
doc = docx.Document(doc_path)

sit_expected = {
    '1.1': 'Berhasil membuka aplikasi dan menampilkan halaman login.',
    '1.2': 'Berhasil melakukan login dan diarahkan ke halaman 2 Factor Authentication.',
    '1.3': 'Menampilkan pesan error email tidak terdaftar.',
    '1.4': 'Menampilkan pesan error autentikasi.',
    '1.5': 'Menampilkan pesan error field mandatory.',
    '1.6': 'Menampilkan pesan error field mandatory.',
    '1.7': 'Menampilkan pesan error field mandatory.'
}

completed_tcs = {'1.1', '1.2', '1.3'}

updated = 0
for tbl in doc.tables:
    r0_c0 = tbl.rows[0].cells[0].text.strip()
    tc_id = None
    is_first_tc = False
    
    if r0_c0 in sit_expected:
        tc_id = r0_c0
        is_first_tc = False
    elif r0_c0 == 'No.' and len(tbl.rows) > 1 and tbl.rows[1].cells[0].text.strip() in sit_expected:
        tc_id = tbl.rows[1].cells[0].text.strip()
        is_first_tc = True
        
    if tc_id:
        exp_cell = tbl.rows[3].cells[1] if is_first_tc else tbl.rows[2].cells[1]
        
        # Clear paragraphs
        for p in list(exp_cell.paragraphs):
            p._element.getparent().remove(p._element)
            
        p = exp_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(5.8)
        
        # Header
        run_hdr = p.add_run("Hasil yang diharapkan [STATUS]:\n")
        run_hdr.font.name = 'Arial'
        run_hdr.font.size = Pt(10)
        run_hdr.bold = True
        
        # SIT Expected Text
        run_text = p.add_run(sit_expected[tc_id])
        run_text.font.name = 'Arial'
        run_text.font.size = Pt(10)
        
        # Status if completed
        if tc_id in completed_tcs:
            run_status = p.add_run(" - [Success]")
            run_status.font.name = 'Arial'
            run_status.font.size = Pt(10)
            run_status.bold = True
            run_status.font.color.rgb = RGBColor(0x2C, 0x52, 0x93)
            
        print(f"Updated TC {tc_id}: {sit_expected[tc_id]} {'[Success]' if tc_id in completed_tcs else ''}")
        updated += 1

print(f"Total updated: {updated}")
doc.save(doc_path)
print("Saved to Drive H.")

for dst in [
    r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
]:
    if os.path.exists(os.path.dirname(dst)):
        shutil.copy2(doc_path, dst)
        print(f"Synced to: {dst}")
