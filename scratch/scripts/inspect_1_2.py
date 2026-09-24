import docx

doc_path = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
doc = docx.Document(doc_path)

for tbl in doc.tables:
    if tbl.rows[0].cells[0].text.strip() == '1.2':
        exp_cell = tbl.rows[2].cells[1]
        for p_idx, p in enumerate(exp_cell.paragraphs):
            print(f"P{p_idx}: {repr(p.text)}")
        break
