import os
import re
import copy
import shutil
import openpyxl
from docx import Document
from docx.table import Table
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from glob import glob

try:
    import win32com.client
    HAS_WIN32 = True
except ImportError:
    HAS_WIN32 = False

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, 'SIT', 'Test Case.xlsx')
if not os.path.exists(EXCEL_PATH) and os.path.exists(r'd:\Project\BTN\SIT\Test Case.xlsx'):
    EXCEL_PATH = r'd:\Project\BTN\SIT\Test Case.xlsx'

template_candidates = [
    os.path.join(BASE_DIR, 'Hasil Uji', 'Template Dokumen Hasil Uji.docx'),
    os.path.join(BASE_DIR, 'Hasil Uji', 'Dokumen_Hasil_Uji_-_UT_Corporate_Banking_CBD.docx'),
    r'd:\Project\BTN\Hasil Uji\Dokumen_Hasil_Uji_-_UT_Corporate_Banking_CBD.docx'
]
TEMPLATE_DOCX = next((p for p in template_candidates if os.path.exists(p)), template_candidates[0])

SCREENSHOT_BASE = os.path.join(BASE_DIR, 'Hasil Uji', 'Screenshot')
OUTPUT_WEB = os.path.join(BASE_DIR, 'Hasil Uji', 'Dokumen_Hasil_Uji_Web.docx')
OUTPUT_MOBILE = os.path.join(BASE_DIR, 'Hasil Uji', 'Dokumen_Hasil_Uji_Mobile.docx')

def sanitize_folder_name(name):
    clean = re.sub(r'[<>:"/\\|?*]', '-', str(name))
    clean = re.sub(r'\s+', ' ', clean).strip(' .')
    if len(clean) > 100:
        clean = clean[:100].strip(' .')
    return clean

def make_long_path(p):
    abs_p = os.path.abspath(p)
    if os.name == 'nt' and not abs_p.startswith('\\\\?\\'):
        return '\\\\?\\' + abs_p
    return abs_p

def clear_cell(cell):
    # Remove all paragraphs
    for p in cell.paragraphs:
        p_element = p._element
        p_element.getparent().remove(p_element)
    # Add one empty paragraph
    cell.add_paragraph()

def set_cell_text(cell, text, bold=False, color_hex=None, alignment=None, space_before=None):
    p = cell.paragraphs[0]
    if alignment is not None:
        p.alignment = alignment
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.bold = bold
    
    if color_hex:
        # docx RGBColor doesn't apply directly to all versions cleanly, so we use XML
        rPr = run._element.get_or_add_rPr()
        color_elm = OxmlElement('w:color')
        color_elm.set(qn('w:val'), color_hex)
        rPr.append(color_elm)
    return p

def set_expected_result(cell, expected_text):
    # Paragraf 1: Hasil yang diharapkan [STATUS]:
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(5.8)
    run1 = p1.add_run("Hasil yang diharapkan [STATUS]:")
    run1.font.name = 'Arial'
    run1.font.size = Pt(10)
    run1.bold = True
    
    # Paragraf 2: [Expected Result] - [Success]
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(6.5)
    
    lines = str(expected_text).strip().split('\n')
    for i, line in enumerate(lines):
        run2 = p2.add_run(line.strip())
        run2.font.name = 'Arial'
        run2.font.size = Pt(10)
        # Set color to 2C5293
        rPr = run2._element.get_or_add_rPr()
        color_elm = OxmlElement('w:color')
        color_elm.set(qn('w:val'), '2C5293')
        rPr.append(color_elm)
        
        if i < len(lines) - 1:
            run_br = p2.add_run('\n')
            
    # Add Success suffix
    run3 = p2.add_run(" - [Success]")
    run3.font.name = 'Arial'
    run3.font.size = Pt(10)
    rPr = run3._element.get_or_add_rPr()
    color_elm = OxmlElement('w:color')
    color_elm.set(qn('w:val'), '2C5293')
    rPr.append(color_elm)


def create_document_for_platform(sheet_name, platform_folder, output_path):
    print(f"Generating for {platform_folder}...")
    # Load original template
    doc = Document(TEMPLATE_DOCX)
    
    # Extract reference tables (Table 4 for first TC, Table 5 for next TCs)
    ref_table_first = copy.deepcopy(doc.tables[4]._tbl)
    ref_table_next = copy.deepcopy(doc.tables[5]._tbl)
    
    # Find the starting point of test cases to delete everything after it
    # We look for "Modul Departemen & RM" (Heading 1)
    body = doc.element.body
    start_delete_idx = -1
    ref_heading_p = None
    
    for i, el in enumerate(body):
        if el.tag == qn('w:p'):
            text = ''.join(t.text or '' for t in el.findall('.//' + qn('w:t')))
            if 'Departemen' in text:
                start_delete_idx = i
                ref_heading_p = copy.deepcopy(el)
                break
                
    if start_delete_idx != -1:
        # Delete from start_delete_idx to the end
        for el in list(body)[start_delete_idx:]:
            body.remove(el)
            
    # Load Excel Data
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[sheet_name]
    
    current_sec = None
    platform_screenshots_dir = os.path.join(SCREENSHOT_BASE, platform_folder)
    drive_h_screenshots_dir = os.path.join(r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot', platform_folder)
    
    for r in range(2, ws.max_row + 1):
        no_val = ws.cell(row=r, column=1).value
        mod_val = ws.cell(row=r, column=2).value
        sub_val = ws.cell(row=r, column=3).value
        title_val = ws.cell(row=r, column=4).value # Test Case Title
        expected_val = ws.cell(row=r, column=8).value # Expected Results
        
        if not mod_val or str(mod_val).strip() == '' or not no_val:
            continue
            
        no_str = str(no_val).strip()
        sec_num = int(no_str.split('.')[0])
        mod_str = str(mod_val).strip()
        sub_str = str(sub_val).strip() if sub_val else ''
        sub_label = f"{mod_str} - {sub_str}" if sub_str else mod_str
        
        tc_title = str(title_val).strip() if title_val else 'Untitled'
        expected_str = str(expected_val).strip() if expected_val else ''
        
        # New submenu section -> Add Heading 1
        if sec_num != current_sec:
            current_sec = sec_num
            
            new_p = copy.deepcopy(ref_heading_p)
            for child in list(new_p):
                if child.tag != qn('w:pPr'):
                    new_p.remove(child)
            
            # Remove auto-numbering to prevent double numbers and weird tab spacing
            pPr = new_p.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    pPr.remove(numPr)
            
            r_run = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            rPr.append(OxmlElement('w:b'))
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '32') # 16pt
            rPr.append(sz)
            szCs = OxmlElement('w:szCs')
            szCs.set(qn('w:val'), '32')
            rPr.append(szCs)
            r_run.append(rPr)
            
            t = OxmlElement('w:t')
            t.text = f"{sec_num}. Modul {sub_label}"
            r_run.append(t)
            new_p.append(r_run)
            
            body.append(new_p)
            body.append(OxmlElement('w:p')) # empty paragraph
            is_first_tc = True
        else:
            is_first_tc = False
            
        # Decide which table template to use
        if is_first_tc:
            new_tbl = copy.deepcopy(ref_table_first)
            t = Table(new_tbl, doc)
            
            # Row 1 (Judul)
            clear_cell(t.rows[1].cells[0])
            set_cell_text(t.rows[1].cells[0], str(no_val), space_before=5.8, alignment=1) # 1=CENTER
            
            clear_cell(t.rows[1].cells[1])
            set_cell_text(t.rows[1].cells[1], f" {tc_title}", space_before=5.8)
            
            image_cell = t.rows[2].cells[1]
            expected_cell = t.rows[3].cells[1]
        else:
            new_tbl = copy.deepcopy(ref_table_next)
            t = Table(new_tbl, doc)
            
            # Row 0 (Judul)
            clear_cell(t.rows[0].cells[0])
            set_cell_text(t.rows[0].cells[0], str(no_val), space_before=5.8, alignment=1)
            
            clear_cell(t.rows[0].cells[1])
            set_cell_text(t.rows[0].cells[1], f" {tc_title}", space_before=5.8)
            
            image_cell = t.rows[1].cells[1]
            expected_cell = t.rows[2].cells[1]
            
        # Set Expected Result
        clear_cell(expected_cell)
        set_expected_result(expected_cell, expected_str)
        
        # Set Image if exists
        clear_cell(image_cell)
        image_p = image_cell.paragraphs[0]
        image_p.alignment = 1 # CENTER
        image_p.paragraph_format.space_before = Pt(5.8)
        
        # Find images in folder
        sub_folder_name = f"{sec_num:02d}. {sanitize_folder_name(sub_label)}"
        tc_folder_name = f"{no_val} {sanitize_folder_name(tc_title)}"
        tc_dir = os.path.join(platform_screenshots_dir, sub_folder_name, tc_folder_name)
        long_tc_dir = make_long_path(tc_dir)
        
        # Fallback to Drive H if local doesn't exist
        if not os.path.exists(long_tc_dir):
            alt_dir = os.path.join(drive_h_screenshots_dir, sub_folder_name, tc_folder_name)
            if os.path.exists(make_long_path(alt_dir)):
                tc_dir = alt_dir
                long_tc_dir = make_long_path(alt_dir)
        
        images_inserted = 0
        if os.path.exists(long_tc_dir):
            valid_exts = ('.png', '.jpg', '.jpeg')
            raw_files = [f for f in os.listdir(long_tc_dir) if f.lower().endswith(valid_exts)]
            files = sorted(raw_files, key=lambda x: [int(c) if c.isdigit() else c for c in re.split(r'(\d+)', x)])
            for f in files:
                img_path = make_long_path(os.path.join(tc_dir, f))
                try:
                    run = image_p.add_run()
                    run.add_picture(img_path, width=Inches(5.9)) # ~15.0 cm
                    run.add_break()
                    images_inserted += 1
                except Exception as e:
                    print(f"Error inserting image {img_path}: {e}")
                    
        # Append table to body after it is completely populated
        body.append(new_tbl)
            
    # Fix TOC field code to use outline levels instead of exact style names (which breaks on localized Word)
    for el in doc.element.body:
        if el.tag == qn('w:sdt'):
            instr = el.findall('.//' + qn('w:instrText'))
            for i in instr:
                if 'TOC' in i.text:
                    i.text = ' TOC \\o "1-3" \\h \\z \\u '
                    
    # Strip native auto-numbering from all Heading paragraphs so our manual numbers work cleanly
    for p in doc.paragraphs:
        if p.style and p.style.name and p.style.name.startswith('Heading'):
            pPr = p._p.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    pPr.remove(numPr)
        
        # Make sure 'Histori Perbaikan' is bold like the other modules
        if 'Histori Perbaikan' in p.text:
            for run in p.runs:
                run.bold = True

    doc.save(output_path)
    print(f"Saved: {output_path}")


def post_process_header(output_path, template_path):
    """
    Inject template's page header (BTN/ZSM logo) and fix page margins
    into the generated DOCX via zipfile manipulation.
    Copies:
      - word/header1.xml          (header content with logos)
      - word/_rels/header1.xml.rels (image references for header)
      - word/media/image278.jpg   (BTN logo referenced by header)
    Patches:
      - word/document.xml         (replace final sectPr with template's sectPr)
      - word/_rels/document.xml.rels (add header1.xml relationship)
    """
    import zipfile, re, shutil, os, tempfile

    print("Injecting page header and fixing margins...")

    # --- Read needed parts from template ---
    with zipfile.ZipFile(template_path, 'r') as zt:
        header_xml   = zt.read('word/header1.xml')
        header_rels  = zt.read('word/_rels/header1.xml.rels')
        btn_logo     = zt.read('word/media/image278.jpg')
        tmpl_doc_rels = zt.read('word/_rels/document.xml.rels').decode('utf-8')
        tmpl_doc_xml  = zt.read('word/document.xml').decode('utf-8')

    # Extract template's LAST sectPr (the one that has headerReference + correct margins)
    all_sectp = re.findall(r'<w:sectPr\b.*?</w:sectPr>', tmpl_doc_xml, re.DOTALL)
    if not all_sectp:
        print("WARNING: Could not find sectPr in template, skipping margin fix.")
        tmpl_sectpr_xml = None
    else:
        tmpl_sectpr_xml = all_sectp[-1]  # Last sectPr has headerReference

    # Extract the header relationship entry from template rels
    # rId285 -> header1.xml
    hdr_rel_match = re.search(
        r'<Relationship[^>]*relationships/header[^>]*/>', tmpl_doc_rels
    )
    hdr_rel_entry = hdr_rel_match.group(0) if hdr_rel_match else None

    # Files we'll add/overwrite at the end (don't copy from input)
    SKIP_FROM_INPUT = {
        'word/header1.xml',
        'word/_rels/header1.xml.rels',
    }

    # --- Patch the output DOCX ---
    tmp_path = output_path + '.tmp'
    with zipfile.ZipFile(output_path, 'r') as zin, \
         zipfile.ZipFile(tmp_path, 'w', compression=zipfile.ZIP_DEFLATED) as zout:

        # Determine a safe new rId for the header relationship in output
        out_doc_rels = zin.read('word/_rels/document.xml.rels').decode('utf-8')
        existing_ids = re.findall(r'Id="(rId\d+)"', out_doc_rels)
        max_id = max((int(i[3:]) for i in existing_ids), default=0)
        new_hdr_rid = f"rId{max_id + 1}"

        # Check if header rel already exists in output
        hdr_already_linked = bool(re.search(r'relationships/header', out_doc_rels))

        # Track which files are in input (so we know what to add fresh)
        input_names = set(zin.namelist())

        for item in zin.infolist():
            # Skip files we'll re-add fresh to avoid duplicate zip entries
            if item.filename in SKIP_FROM_INPUT:
                continue

            data = zin.read(item.filename)

            if item.filename == '[Content_Types].xml':
                # Register header1.xml content type
                ct_str = data.decode('utf-8')
                hdr_ct = (
                    '<Override PartName="/word/header1.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>'
                )
                if 'header1.xml' not in ct_str:
                    ct_str = ct_str.replace('</Types>', hdr_ct + '</Types>')
                data = ct_str.encode('utf-8')

            elif item.filename == 'word/_rels/document.xml.rels':
                # Add header relationship if not already there
                if not hdr_already_linked:
                    rel_str = data.decode('utf-8')
                    new_rel = (
                        f'<Relationship Id="{new_hdr_rid}" '
                        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" '
                        f'Target="header1.xml"/>'
                    )
                    rel_str = rel_str.replace('</Relationships>', new_rel + '</Relationships>')
                    data = rel_str.encode('utf-8')
                # If already linked, keep as-is and use existing rId
                else:
                    # Find existing header rId
                    m = re.search(r'Id="(rId\d+)"[^>]*relationships/header', data.decode('utf-8'))
                    if not m:
                        m = re.search(r'relationships/header[^>]*Id="(rId\d+)"', data.decode('utf-8'))
                    if m:
                        new_hdr_rid = m.group(1)

            elif item.filename == 'word/document.xml':
                # Replace final sectPr with template's (has headerReference + margins)
                if tmpl_sectpr_xml:
                    doc_str = data.decode('utf-8')
                    # Start with template's sectPr
                    patched_sectpr = tmpl_sectpr_xml
                    # Remove footerReference (we don't have footer1.xml in output)
                    patched_sectpr = re.sub(
                        r'<w:footerReference[^/]*/>', '', patched_sectpr
                    )
                    # Update headerReference rId to match new_hdr_rid
                    patched_sectpr = re.sub(
                        r'(<w:headerReference[^>]*)r:id="rId\d+"',
                        rf'\g<1>r:id="{new_hdr_rid}"',
                        patched_sectpr
                    )
                    # Replace last sectPr in output safely
                    last_sect_idx = doc_str.rfind('<w:sectPr')
                    if last_sect_idx != -1:
                        end_sect_idx = doc_str.find('</w:sectPr>', last_sect_idx) + len('</w:sectPr>')
                        doc_str = doc_str[:last_sect_idx] + patched_sectpr + doc_str[end_sect_idx:]
                    data = doc_str.encode('utf-8')

            zout.writestr(item, data)

        # Add header files fresh (these were skipped from input to avoid duplicates)
        zout.writestr('word/header1.xml', header_xml)
        zout.writestr('word/_rels/header1.xml.rels', header_rels)
        # Add BTN logo if not already in output
        if 'word/media/image278.jpg' not in input_names:
            zout.writestr('word/media/image278.jpg', btn_logo)

    # Replace original with patched
    os.replace(tmp_path, output_path)
    print("Page header injected successfully.")


def update_toc(doc_path):
    if not HAS_WIN32:
        print("win32com not available, skipping automatic TOC update.")
        return
    try:
        print(f"Updating TOC for {doc_path} ...")
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(doc_path), ConfirmConversions=False, ReadOnly=False)
        
        # Ensure TOC exists and update it
        if doc.TablesOfContents.Count > 0:
            doc.TablesOfContents(1).Update()
            
        doc.Close(SaveChanges=True)
        word.Quit()
        print("TOC successfully updated!")
    except Exception as e:
        print(f"Failed to update TOC automatically: {e}")
        try:
            word.Quit()
        except:
            pass

def main():
    create_document_for_platform('TC BTN SMART Web', 'Web', OUTPUT_WEB)
    update_toc(OUTPUT_WEB)
    post_process_header(OUTPUT_WEB, TEMPLATE_DOCX)
    
    drive_output_web = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Web.docx'
    os.makedirs(os.path.dirname(drive_output_web), exist_ok=True)
    shutil.copy2(OUTPUT_WEB, drive_output_web)
    print(f"Synced to Drive H: {drive_output_web}")
    
    drive_g_output = r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Web.docx'
    os.makedirs(os.path.dirname(drive_g_output), exist_ok=True)
    shutil.copy2(OUTPUT_WEB, drive_g_output)
    print(f"Synced to Drive G: {drive_g_output}")
    
    print("Done!")

if __name__ == '__main__':
    main()
