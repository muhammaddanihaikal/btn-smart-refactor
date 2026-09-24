import os
import re
import shutil
import openpyxl

EXCEL_PATH = r'd:\Project\BTN\SIT\Test Case.xlsx'
BACKUP_DIR = r'd:\Project\BTN\scratch\all_images_backup'

TARGET_DIRS = [
    r'd:\Project\BTN\Hasil Uji\Screenshot\Web',
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Web',
]
if os.path.exists(r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Web'):
    TARGET_DIRS.append(r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Web')

def sanitize_folder_name(name):
    clean = re.sub(r'[<>:"/\\|?*]', '-', str(name))
    clean = re.sub(r'\s+', ' ', clean).strip(' .')
    if len(clean) > 90:
        clean = clean[:90].strip(' .')
    return clean

def make_long_path(p):
    abs_p = os.path.abspath(p)
    if os.name == 'nt' and not abs_p.startswith('\\\\?\\'):
        return '\\\\?\\' + abs_p
    return abs_p

# Load Excel test cases
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['TC BTN SMART Web']

tc_list = []
for r in range(2, ws.max_row + 1):
    no_val = ws.cell(r, 1).value
    mod_val = ws.cell(r, 2).value
    sub_val = ws.cell(r, 3).value
    title_val = ws.cell(r, 4).value
    desc_val = ws.cell(r, 5).value
    prio_val = ws.cell(r, 6).value
    step_val = ws.cell(r, 7).value
    exp_val = ws.cell(r, 8).value
    
    if mod_val and str(mod_val).strip() and title_val and str(title_val).strip():
        tc_list.append({
            'no': str(no_val).strip(),
            'module': str(mod_val).strip(),
            'submenu': str(sub_val).strip() if sub_val else '',
            'title': str(title_val).strip(),
            'desc': str(desc_val).strip() if desc_val else '',
            'prio': str(prio_val).strip() if prio_val else '',
            'steps': str(step_val).strip() if step_val else '-',
            'expected': str(exp_val).strip() if exp_val else '-'
        })

print(f"Total TCs loaded from Excel: {len(tc_list)}")

# Mapping from backup image paths to new TC numbers
# Old 4.1 -> 6.1 (Membuka halaman User)
# Old 4.2 -> 6.2 (Mencari data User dengan keyword valid)
# Old 4.4 -> 6.4 (Menambahkan data User)
image_mappings = {
    '4.1': ('6.1', ['1.png']),
    '4.2': ('6.2', ['1.png']),
    '4.4': ('6.4', ['1.png', '2.png', '3.png']),
}

for base_dir in TARGET_DIRS:
    print(f"\nProcessing target directory: {base_dir}")
    os.makedirs(base_dir, exist_ok=True)
    
    # Keep track of created new sub menu folders
    created_submenus = set()
    tc_folder_map = {} # no -> full path of tc folder
    
    for tc in tc_list:
        no_val = tc['no']
        sec_num = int(no_val.split('.')[0])
        mod_name = tc['module']
        sub_name = tc['submenu']
        
        sub_label = f"{mod_name} - {sub_name}" if sub_name else mod_name
        folder_sub = f"{sec_num:02d}. {sanitize_folder_name(sub_label)}"
        folder_tc = f"{no_val} {sanitize_folder_name(tc['title'])}"
        
        sub_dir = os.path.join(base_dir, folder_sub)
        tc_dir = os.path.join(sub_dir, folder_tc)
        
        long_tc_dir = make_long_path(tc_dir)
        os.makedirs(long_tc_dir, exist_ok=True)
        created_submenus.add(folder_sub)
        tc_folder_map[no_val] = tc_dir
        
        # Write .txt info file
        txt_name = f"{no_val} {sanitize_folder_name(tc['title'])}.txt"
        txt_path = make_long_path(os.path.join(tc_dir, txt_name))
        
        content = [
            f"NO TEST CASE : {no_val}",
            f"JUDUL        : {tc['title']}",
            f"MODUL        : {mod_name}",
        ]
        if sub_name:
            content.append(f"SUB MENU     : {sub_name}")
        if tc['desc']:
            content.append(f"DESKRIPSI    : {tc['desc']}")
        if tc['prio']:
            content.append(f"PRIORITAS    : {tc['prio']}")
            
        content.append("\n" + "=" * 50)
        content.append("STEPS (LANGKAH-LANGKAH PENGUJIAN):")
        content.append("=" * 50)
        content.append(tc['steps'])
        
        content.append("\n" + "=" * 50)
        content.append("EXPECTED RESULTS (HASIL YANG DIHARAPKAN):")
        content.append("=" * 50)
        content.append(tc['expected'])
        content.append("")
        
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(content))

    print(f"Created/verified {len(tc_list)} TC folders across {len(created_submenus)} submenus in {base_dir}")
    
    # Copy backup images into new folders
    copied_images = 0
    for old_no, (new_no, img_list) in image_mappings.items():
        dst_folder = tc_folder_map.get(new_no)
        if not dst_folder:
            print(f"WARNING: Target folder for {new_no} not found!")
            continue
            
        # Find matching old folder in backup
        for old_tc_folder in os.listdir(os.path.join(BACKUP_DIR, '04. User Authority')):
            if old_tc_folder.startswith(old_no + ' '):
                src_folder = os.path.join(BACKUP_DIR, '04. User Authority', old_tc_folder)
                for img_name in img_list:
                    src_img = os.path.join(src_folder, img_name)
                    dst_img = make_long_path(os.path.join(dst_folder, img_name))
                    if os.path.exists(src_img):
                        shutil.copy2(src_img, dst_img)
                        copied_images += 1
                        print(f"Copied image: {src_img} -> {dst_img}")
    
    print(f"Total images restored in {base_dir}: {copied_images}")
    
    # Remove old 20 module top-level folders (e.g. '01. Login', '04. User Authority')
    # BUT only if they are not in created_submenus!
    # e.g. '01. Login' matches created_submenus, but '04. User Authority' does not (now '06. User Authority - User', etc.)
    all_top = os.listdir(base_dir)
    for item in all_top:
        if item not in created_submenus:
            item_path = make_long_path(os.path.join(base_dir, item))
            if os.path.isdir(item_path):
                print(f"Removing obsolete folder: {item}")
                shutil.rmtree(item_path)

print("\n--- ALL TARGET DIRECTORIES SYNCHRONIZED SUCCESSFULLY ---")
