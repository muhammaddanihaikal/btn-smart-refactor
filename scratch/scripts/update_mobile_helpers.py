import os
import re
import openpyxl
import docx

# Paths
EXCEL_PATH = r'D:\Project\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx'
DOCX_PATH  = r'D:\Project\BTN Smart\Refactor\SIT\Document SIT BTN Smart\SIT BTN SMART Mobile - 24 Sep 2026.docx'
LOCAL_DIR  = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile'
H_DIR      = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile'
G_DIR      = r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile'

def sanitize_folder_name(name):
    clean = re.sub(r'[<>:"/\\|?*]', '-', str(name))
    clean = re.sub(r'\s+', ' ', clean).strip(' .')
    if len(clean) > 100:
        clean = clean[:100].strip(' .')
    return clean

# 1. Load Excel Data
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['TC BTN SMART Mobile']

excel_data = {}
for r in range(2, ws.max_row + 1):
    title = ws.cell(r, 3).value
    if title:
        t_clean = str(title).strip().lower()
        excel_data[t_clean] = {
            'mod': str(ws.cell(r, 1).value or '').strip(),
            'sub': str(ws.cell(r, 2).value or '').strip(),
            'title': str(title).strip(),
            'desc': str(ws.cell(r, 4).value or 'Positive Case').strip(),
            'prio': str(ws.cell(r, 5).value or 'High').strip(),
            'steps': str(ws.cell(r, 6).value or '').strip(),
            'expected': str(ws.cell(r, 7).value or '').strip()
        }

print(f"Loaded {len(excel_data)} TCs from Excel.")

# 2. Load Docx TCs
doc = docx.Document(DOCX_PATH)
tcs = []
current_mod = ''
mod_index = 0
last_mod = ''

for child in doc.element.body:
    if child.tag.endswith('p'):
        p = docx.text.paragraph.Paragraph(child, doc)
        txt = p.text.strip()
        if 'Modul ' in txt and txt[0].isdigit():
            current_mod = txt.split('Modul ')[-1].strip()
            if current_mod != last_mod:
                mod_index += 1
                last_mod = current_mod
    elif child.tag.endswith('tbl'):
        tbl = docx.table.Table(child, doc)
        if len(tbl.columns) == 7:
            for r in tbl.rows:
                c0 = r.cells[0].text.strip()
                if c0 and '.' in c0 and c0[0].isdigit():
                    tcs.append({
                        'mod_idx': mod_index,
                        'mod_name': current_mod,
                        'no': c0,
                        'title': r.cells[1].text.strip(),
                        'scenario': r.cells[2].text.strip(),
                        'expected': r.cells[3].text.strip()
                    })

print(f"Extracted {len(tcs)} TCs from SIT Mobile Doc.")

def format_numbered(text):
    text = text.strip()
    if not text:
        return ""
    # If already starts with number e.g. "1."
    if re.match(r'^\d+\.', text):
        return text
    # Split by any dash separator ' - ', ' \u2013 ', ' \u2014 ', or '\n'
    # normalize various dashes
    normalized = re.sub(r'\s+[\-\u2013\u2014]\s+', ' \n ', text)
    if '\n' in normalized:
        parts = [p.strip() for p in normalized.split('\n') if p.strip()]
    else:
        parts = [text]
    
    numbered = []
    for i, p in enumerate(parts, 1):
        clean_p = p.lstrip('- ').strip()
        if not clean_p.endswith('.'):
            clean_p += '.'
        numbered.append(f"{i}. {clean_p}")
    return "\n".join(numbered)

# 3. Process and write .txt files
def update_directory(target_base):
    if not os.path.exists(target_base):
        print(f"Directory not found: {target_base}")
        return
    
    long_prefix = "\\\\?\\"
    abs_base = long_prefix + os.path.abspath(target_base)
    
    updated_count = 0
    for tc in tcs:
        t_key = tc['title'].strip().lower()
        ex = excel_data.get(t_key)
        
        if ex and ex['steps']:
            steps_text = format_numbered(ex['steps'])
            exp_text = format_numbered(ex['expected'])
            desc_text = ex['desc']
            prio_text = ex['prio']
        else:
            steps_text = format_numbered(tc['scenario'])
            exp_text = format_numbered(tc['expected'])
            is_neg = any(k in tc['title'].lower() or k in tc['scenario'].lower() for k in ['tidak', 'salah', 'negative', 'negatve', 'gagal'])
            desc_text = 'Negative Case' if is_neg else 'Positive Case'
            prio_text = 'High'
            
        folder_sub = f"{tc['mod_idx']:02d}. {sanitize_folder_name(tc['mod_name'])}"
        folder_tc = f"{tc['no']} {sanitize_folder_name(tc['title'])}"
        dir_path = os.path.join(abs_base, folder_sub, folder_tc)
        
        if not os.path.exists(dir_path):
            os.makedirs(dir_path, exist_ok=True)
            
        txt_path = os.path.join(dir_path, f"{folder_tc}.txt")
        
        file_content = f"""NO TEST CASE : {tc['no']}
JUDUL        : {tc['title']}
MODUL        : {tc['mod_name']}
DESKRIPSI    : {desc_text}
PRIORITAS    : {prio_text}

==================================================
STEPS (LANGKAH-LANGKAH PENGUJIAN):
==================================================
{steps_text}

==================================================
EXPECTED RESULTS (HASIL YANG DIHARAPKAN):
==================================================
{exp_text}
"""
        with open(txt_path, 'w', encoding='utf-8') as fp:
            fp.write(file_content)
        updated_count += 1

    print(f"Updated {updated_count} files in {target_base}")

print("Updating Local...")
update_directory(LOCAL_DIR)
print("Updating Drive H...")
update_directory(H_DIR)
print("Updating Drive G...")
update_directory(G_DIR)
print("ALL DONE!")
