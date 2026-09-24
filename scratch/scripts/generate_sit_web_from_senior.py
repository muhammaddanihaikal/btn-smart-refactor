import os, copy, zipfile, shutil, re
import docx
import openpyxl
from collections import OrderedDict
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SENIOR_DOCX   = r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Web (Updated).docx'
TEMPLATE_PATH = r'D:\Project\BTN Smart\Refactor\SIT\Form Script - Skenario SIT BTN SMART Upgrade Server.docx'
UJI_SISTEM    = r'D:\Project\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx'
OUTPUT_WEB    = r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx'
OUTPUT_WEB_H  = r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx'
OUTPUT_WEB_G  = r'G:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx'

print(f"Loading mapping from {UJI_SISTEM} ...")
wb = openpyxl.load_workbook(UJI_SISTEM, data_only=True)
ws = wb['TC BTN SMART Web']
excel_map = {}
for r in range(2, ws.max_row + 1):
    mod = str(ws.cell(r, 1).value or '').strip()
    sub = str(ws.cell(r, 2).value or '').strip()
    title = str(ws.cell(r, 3).value or '').strip()
    if title:
        if title.lower() not in excel_map:
            excel_map[title.lower()] = []
        excel_map[title.lower()].append((mod, sub))

# Manual fallbacks
manual_map = {
    'menginput data step 4': [('Lead Generation', 'Prospek Individu')],
    'menginput data step 3': [('Lead Generation', 'Prospek Lembaga')],
    'melakukan aktifitas marketing (call)': [('Lead Qualification', 'Aktifitas')],
    'melakukan aktifitas marketing prospek (call) dengan status \'nasabah menolak\'': [('Lead Qualification', 'Aktifitas')],
    'validasi data visit in prospek': [('Lead Qualification', 'Aktifitas')],
    'melakukan aktifitas marketing open account': [('Lead Qualification', 'Aktifitas')],
}
excel_map.update(manual_map)

print(f"Extracting TCs from senior doc: {SENIOR_DOCX} ...")
doc_senior = docx.Document(SENIOR_DOCX)
final_submenus = OrderedDict()
total_tcs = 0

current_mod = 'Unknown'
for child in list(doc_senior.element.body):
    if child.tag.endswith('p'):
        p = docx.text.paragraph.Paragraph(child, doc_senior)
        txt = p.text.strip()
        if txt.startswith('Modul ') or p.style.name.startswith('Heading'):
            current_mod = txt.replace('Modul ', '', 1).strip()
    elif child.tag.endswith('tbl'):
        tbl = docx.table.Table(child, doc_senior)
        if len(tbl.columns) == 7 and len(tbl.rows) > 0:
            c0 = tbl.rows[0].cells[0].text.strip()
            if c0 == 'No' or (c0 and '.' in c0 and c0[0].isdigit()):
                # Parse TCs
                for r in tbl.rows:
                    rc0 = r.cells[0].text.strip()
                    if rc0 and '.' in rc0 and rc0[0].isdigit():
                        tc_title = r.cells[1].text.strip()
                        tc_scen  = r.cells[2].text.strip()
                        tc_exp   = r.cells[3].text.strip()
                        tc_act   = r.cells[4].text.strip()
                        tc_stat  = r.cells[5].text.strip()
                        tc_rem   = r.cells[6].text.strip()
                        
                        possibles = excel_map.get(tc_title.lower(), [('?', '?')])
                        best_match = possibles[0]
                        for p_map in possibles:
                            if current_mod.lower() in p_map[0].lower() or p_map[0].lower() in current_mod.lower():
                                best_match = p_map
                                break
                        
                        # Add to final
                        key = (best_match[0], best_match[1])
                        if key not in final_submenus:
                            final_submenus[key] = []
                        final_submenus[key].append({
                            'title': tc_title,
                            'scenario': tc_scen,
                            'expected': tc_exp,
                            'actual': tc_act,
                            'status': tc_stat,
                            'remarks': tc_rem
                        })
                        total_tcs += 1

print(f"Total active submenus: {len(final_submenus)}")
print(f"Total TCs: {total_tcs}")

# ── XML HELPERS ──
doc_template    = Document(TEMPLATE_PATH)
body_tpl        = list(doc_template.element.body)
TBL_HEADER      = doc_template.tables[0]
TBL_INFO        = doc_template.tables[1]
TBL_DATA_TPL    = doc_template.tables[2]
DAFTAR_TEMPLATE = body_tpl[20]

def set_cell_center(tc_elem, text, bold=False):
    tcPr = tc_elem.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr'); tc_elem.insert(0, tcPr)
    va = tcPr.find(qn('w:vAlign'))
    if va is None:
        va = OxmlElement('w:vAlign'); tcPr.append(va)
    va.set(qn('w:val'), 'center')
    text = text if isinstance(text, str) else (str(text) if text else '')
    lines = text.split('\n')
    paras = tc_elem.findall(qn('w:p'))
    while len(tc_elem.findall(qn('w:p'))) < len(lines):
        p_new = copy.deepcopy(paras[0])
        for ch in list(p_new):
            if ch.tag != qn('w:pPr'): p_new.remove(ch)
        tc_elem.append(p_new)
    paras = tc_elem.findall(qn('w:p'))
    for i, line in enumerate(lines):
        p = paras[i]
        pPr = p.find(qn('w:pPr'))
        if pPr is None: pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
        jc = pPr.find(qn('w:jc'))
        if jc is None: jc = OxmlElement('w:jc'); pPr.append(jc)
        jc.set(qn('w:val'), 'center')
        for ch in list(p):
            if ch.tag not in [qn('w:pPr'), qn('w:pPrChange')]: p.remove(ch)
        r_elem = OxmlElement('w:r')
        if bold:
            rPr = OxmlElement('w:rPr'); rPr.append(OxmlElement('w:b')); r_elem.append(rPr)
        t = OxmlElement('w:t')
        t.text = line
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r_elem.append(t); p.append(r_elem)
    for extra in tc_elem.findall(qn('w:p'))[len(lines):]:
        tc_elem.remove(extra)

def build_info_table(submenus_list, total_tc):
    tbl = copy.deepcopy(TBL_INFO._tbl)
    rows = tbl.findall(qn('w:tr'))
    r0_cells = rows[0].findall(qn('w:tc'))
    set_cell_center(r0_cells[1], 'Jumlah Script : ' + str(total_tc))
    r1c0 = rows[1].findall(qn('w:tc'))[0]
    orig_paras = r1c0.findall(qn('w:p'))
    bullet_pPr = None
    if len(orig_paras) > 1:
        pPr = orig_paras[1].find(qn('w:pPr'))
        if pPr is not None: bullet_pPr = copy.deepcopy(pPr)
    for p in r1c0.findall(qn('w:p')):
        for ch in list(p):
            if ch.tag not in [qn('w:pPr'), qn('w:pPrChange')]: p.remove(ch)
    needed = 1 + len(submenus_list)
    while len(r1c0.findall(qn('w:p'))) < needed:
        r1c0.append(OxmlElement('w:p'))
    paras = r1c0.findall(qn('w:p'))
    p0 = paras[0]
    r0 = OxmlElement('w:r')
    r0Pr = OxmlElement('w:rPr'); r0Pr.append(OxmlElement('w:b'))
    sz0 = OxmlElement('w:sz'); sz0.set(qn('w:val'), '20'); r0Pr.append(sz0)
    r0.append(r0Pr)
    t0 = OxmlElement('w:t'); t0.text = 'Modul :'; r0.append(t0); p0.append(r0)
    for idx, (mod_name, sub_name) in enumerate(submenus_list):
        label = 'Modul ' + mod_name + (' - ' + sub_name if sub_name else '')
        p = paras[1 + idx]
        if bullet_pPr is not None:
            old = p.find(qn('w:pPr'))
            if old is not None: p.remove(old)
            cloned = copy.deepcopy(bullet_pPr)
            sp = cloned.find(qn('w:spacing'))
            if sp is not None: sp.set(qn('w:before'), '20'); sp.set(qn('w:after'), '0')
            p.insert(0, cloned)
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '20'); rPr.append(sz); r.append(rPr)
        t = OxmlElement('w:t'); t.text = label
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t); p.append(r)
    set_cell_center(rows[2].findall(qn('w:tc'))[0], '')
    return tbl

def build_title_para(sec_idx, mod_name, sub_name):
    p = copy.deepcopy(DAFTAR_TEMPLATE)
    pPr = p.find(qn('w:pPr'))
    if pPr is not None:
        for tag in [qn('w:numPr'), qn('w:tabs')]:
            el = pPr.find(tag)
            if el is not None: pPr.remove(el)
        ind = pPr.find(qn('w:ind'))
        if ind is None: ind = OxmlElement('w:ind'); pPr.append(ind)
        ind.set(qn('w:left'), '259'); ind.set(qn('w:hanging'), '0')
        sp = pPr.find(qn('w:spacing'))
        if sp is None: sp = OxmlElement('w:spacing'); pPr.append(sp)
        sp.set(qn('w:before'), '120'); sp.set(qn('w:after'), '160')
        if sec_idx > 1:
            pbb = pPr.find(qn('w:pageBreakBefore'))
            if pbb is None:
                pbb = OxmlElement('w:pageBreakBefore')
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None: pStyle.addnext(pbb)
                else: pPr.insert(0, pbb)
    orig_runs = p.findall(qn('w:r'))
    orig_rpr = None
    if orig_runs:
        rPr_el = orig_runs[0].find(qn('w:rPr'))
        if rPr_el is not None: orig_rpr = copy.deepcopy(rPr_el)
    for ch in list(p):
        if ch.tag not in [qn('w:pPr'), qn('w:pPrChange')]: p.remove(ch)
    label = 'Modul ' + mod_name + (' - ' + sub_name if sub_name else '')
    title = str(sec_idx) + '. ' + label
    r_elem = OxmlElement('w:r')
    if orig_rpr is not None: r_elem.append(orig_rpr)
    t = OxmlElement('w:t'); t.text = title
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r_elem.append(t); p.append(r_elem)
    return p

def zero_height_para(sect_pr=None):
    p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')
    sp.set(qn('w:line'), '1'); sp.set(qn('w:lineRule'), 'exact')
    pPr.append(sp)
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rPr.append(sz); pPr.append(rPr)
    if sect_pr is not None: pPr.append(copy.deepcopy(sect_pr))
    p.append(pPr); return p

def empty_para():
    p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle'); pStyle.set(qn('w:val'), 'TeksIsi')
    pPr.append(pStyle); p.insert(0, pPr); return p

def build_data_table(tcs, sec_idx):
    tbl = copy.deepcopy(TBL_DATA_TPL._tbl)
    rows_elem = tbl.findall(qn('w:tr'))
    for r in rows_elem[1:]: tbl.remove(r)
    header_tr = rows_elem[0]
    for c in header_tr.findall(qn('w:tc')):
        tcPr = c.find(qn('w:tcPr'))
        if tcPr is not None:
            va = tcPr.find(qn('w:vAlign'))
            if va is None: va = OxmlElement('w:vAlign'); tcPr.append(va)
            va.set(qn('w:val'), 'center')
        for p in c.findall(qn('w:p')):
            pPr = p.find(qn('w:pPr'))
            if pPr is None: pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
            jc = pPr.find(qn('w:jc'))
            if jc is None: jc = OxmlElement('w:jc'); pPr.append(jc)
            jc.set(qn('w:val'), 'center')
    tmpl_row = TBL_DATA_TPL.rows[1]
    for i, tc in enumerate(tcs, start=1):
        tc_number = str(sec_idx) + '.' + str(i)
        new_tr = copy.deepcopy(tmpl_row._tr)
        cells = new_tr.findall(qn('w:tc'))
        set_cell_center(cells[0], tc_number)
        set_cell_center(cells[1], tc['title'])
        set_cell_center(cells[2], tc['scenario'])
        set_cell_center(cells[3], tc['expected'])
        set_cell_center(cells[4], tc['actual'])
        set_cell_center(cells[5], tc['status'] or 'P')
        set_cell_center(cells[6], tc['remarks'])
        tbl.append(new_tr)
    return tbl


def strip_para_ids(docx_path):
    tmp = docx_path + '.tmp_para'
    with zipfile.ZipFile(docx_path, 'r') as z_in, zipfile.ZipFile(tmp, 'w') as z_out:
        for item in z_in.infolist():
            if item.filename == 'word/document.xml':
                content = z_in.read(item.filename)
                try:
                    import lxml.etree as etree
                    root = etree.fromstring(content)
                    for el in root.iter():
                        for k in list(el.attrib.keys()):
                            if 'paraId' in k or 'textId' in k:
                                del el.attrib[k]
                    content = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone='yes')
                except: pass
                z_out.writestr(item, content)
            else:
                z_out.writestr(item, z_in.read(item.filename))
    import shutil
    shutil.move(tmp, docx_path)

def post_process(docx_path):
    tmp = docx_path + '.tmp_pp'
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    with zipfile.ZipFile(TEMPLATE_PATH, 'r') as z_src:
        template_header = (z_src.read('word/header1.xml').decode('utf-8') if 'word/header1.xml' in z_src.namelist() else None)
    if template_header:
        page_idx = template_header.find(' PAGE </w:instrText>')
        if page_idx != -1:
            ac_start = template_header.rfind('<mc:AlternateContent>', 0, page_idx)
            ac_end = template_header.find('</mc:AlternateContent>', page_idx) + len('</mc:AlternateContent>')
            r_start = template_header.rfind('<w:r ', 0, ac_start)
            r_end = template_header.find('</w:r>', ac_end) + len('</w:r>')
            template_header = template_header[:r_start] + template_header[r_end:]
            inline_page = (
                '<w:p><w:pPr><w:pStyle w:val="TeksIsi"/><w:jc w:val="right"/>'
                '<w:spacing w:before="530"/><w:ind w:right="300"/>'
                '<w:rPr><w:b/><w:sz w:val="20"/><w:szCs w:val="20"/>'
                '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr></w:pPr>'
                '<w:r><w:rPr><w:b/><w:sz w:val="20"/><w:szCs w:val="20"/>'
                '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr>'
                '<w:t xml:space="preserve">Page </w:t></w:r>'
                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                '<w:r><w:rPr><w:b/><w:sz w:val="20"/><w:szCs w:val="20"/>'
                '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr>'
                '<w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                '<w:r><w:rPr><w:b/><w:sz w:val="20"/><w:szCs w:val="20"/>'
                '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr>'
                '<w:t>2</w:t></w:r>'
                '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
            )
            template_header = template_header.replace('</w:hdr>', inline_page + '</w:hdr>')
    with zipfile.ZipFile(docx_path, 'r') as z_in, zipfile.ZipFile(tmp, 'w') as z_out:
        for item in z_in.infolist():
            if item.filename == 'word/header1.xml' and template_header:
                z_out.writestr(item, template_header.encode('utf-8'))
            elif item.filename == 'word/settings.xml':
                content = z_in.read(item.filename)
                try:
                    import lxml.etree as etree
                    root = etree.fromstring(content)
                    for uf in root.findall('.//w:updateFields', namespaces={'w': W}):
                        root.remove(uf)
                    content = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone='yes')
                except: pass
                z_out.writestr(item, content)
            else:
                z_out.writestr(item, z_in.read(item.filename))
    shutil.move(tmp, docx_path)

print(f'\nBuilding document: {OUTPUT_WEB}')
doc = Document(TEMPLATE_PATH)
body = doc.element.body
inner_sect_pr = None
final_sect_pr = None
for child in list(body):
    if child.tag == qn('w:p'):
        pPr = child.find(qn('w:pPr'))
        if pPr is not None:
            sp = pPr.find(qn('w:sectPr'))
            if sp is not None: inner_sect_pr = copy.deepcopy(sp)
    elif child.tag == qn('w:sectPr'):
        final_sect_pr = copy.deepcopy(child)
for child in list(body): body.remove(child)

body.append(copy.deepcopy(TBL_HEADER._tbl))
body.append(empty_para())

# ── Sort final_submenus based on EXCEL order ──
import openpyxl
wb_excel = openpyxl.load_workbook(r'D:\Project\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx', data_only=True)
ws_excel = wb_excel['TC BTN SMART Web']
excel_order = []
for r in range(2, ws_excel.max_row + 1):
    mod = ws_excel.cell(r, 2).value
    sub = ws_excel.cell(r, 3).value
    if mod:
        m = str(mod).strip()
        s = str(sub).strip() if sub else ''
        key = (m, s)
        if not excel_order or excel_order[-1] != key:
            excel_order.append(key)

ordered_keys = []
for k in excel_order:
    if k in final_submenus:
        ordered_keys.append(k)

# Add any keys that were in final_submenus but somehow not in Excel
for k in final_submenus.keys():
    if k not in ordered_keys:
        ordered_keys.append(k)

# Rebuild final_submenus as a new ordered dict
new_final_submenus = {k: final_submenus[k] for k in ordered_keys}
final_submenus = new_final_submenus

body.append(build_info_table(list(final_submenus.keys()), total_tcs))
body.append(zero_height_para(inner_sect_pr))

for sec_idx, ((mod_name, sub_name), tcs) in enumerate(final_submenus.items(), 1):
    body.append(build_title_para(sec_idx, mod_name, sub_name))
    body.append(build_data_table(tcs, sec_idx))

body.append(zero_height_para())
if final_sect_pr is not None:
    pgNum = final_sect_pr.find(qn('w:pgNumType'))
    if pgNum is not None: final_sect_pr.remove(pgNum)
    body.append(final_sect_pr)

doc.save(OUTPUT_WEB)
print('Saved:', OUTPUT_WEB)
strip_para_ids(OUTPUT_WEB)
post_process(OUTPUT_WEB)
print('Post-processed:', OUTPUT_WEB)

for dst in [OUTPUT_WEB_H, OUTPUT_WEB_G]:
    shutil.copy2(OUTPUT_WEB, dst)
    print('Synced to:', dst)
print('\nDone!')
