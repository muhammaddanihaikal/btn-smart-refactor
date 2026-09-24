import os, re, shutil
import openpyxl
import docx

REPLACEMENTS = [
    ("Membuka halaman Login Non Employee", "Membuka halaman Login"),
    ("Melakukan login Non Employee dengan email dan password valid", "Melakukan login dengan email dan password valid"),
    ("Melakukan login Non Employee dengan email tidak terdaftar", "Melakukan login dengan email tidak terdaftar"),
    ("Melakukan login Non Employee dengan password salah", "Melakukan login dengan password salah"),
]

# 1. UPDATE EXCEL FILES
excel_files = [
    r'D:\Project\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx',
    r'D:\Project\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx',
    r'D:\Project\BTN Smart\Refactor\SIT\Test Case.xlsx'
]

for ef in excel_files:
    if os.path.exists(ef):
        wb = openpyxl.load_workbook(ef)
        modified = False
        for sname in wb.sheetnames:
            if 'mobile' in sname.lower() or 'web' in sname.lower():
                ws = wb[sname]
                for r in range(2, 20):
                    for c in range(1, 10):
                        val = ws.cell(r, c).value
                        if val and isinstance(val, str):
                            new_val = val
                            for old_t, new_t in REPLACEMENTS:
                                if old_t in new_val:
                                    new_val = new_val.replace(old_t, new_t)
                            if new_val != val:
                                ws.cell(r, c).value = new_val
                                modified = True
        if modified:
            wb.save(ef)
            print(f"Updated Excel: {ef}")

# 2. UPDATE WORD DOCS
def replace_in_docx(doc_path):
    if not os.path.exists(doc_path):
        return
    doc = docx.Document(doc_path)
    changed = False
    
    # Check paragraphs
    for p in doc.paragraphs:
        for old_t, new_t in REPLACEMENTS:
            if old_t in p.text:
                for run in p.runs:
                    if old_t in run.text:
                        run.text = run.text.replace(old_t, new_t)
                        changed = True
                if old_t in p.text: # fallback if split runs
                    p.text = p.text.replace(old_t, new_t)
                    changed = True
                    
    # Check tables
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for old_t, new_t in REPLACEMENTS:
                    if old_t in cell.text:
                        for p in cell.paragraphs:
                            if old_t in p.text:
                                for r in p.runs:
                                    if old_t in r.text:
                                        r.text = r.text.replace(old_t, new_t)
                                        changed = True
                                if old_t in p.text:
                                    p.text = p.text.replace(old_t, new_t)
                                    changed = True
    if changed:
        doc.save(doc_path)
        print(f"Updated Docx: {doc_path}")

doc_targets = [
    r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile.docx',
    r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile.docx',
    r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx',
    r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web.docx',
    r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx',
]

for dt in doc_targets:
    replace_in_docx(dt)

# 3. RENAME FOLDERS AND TXT FILES IN SCREENSHOT/MOBILE
ss_roots = [
    r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile\01. Login',
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile\01. Login',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Screenshot\Mobile\01. Login',
]

for root_dir in ss_roots:
    if not os.path.exists(root_dir):
        continue
    long_prefix = "\\\\?\\"
    abs_root = long_prefix + os.path.abspath(root_dir)
    
    for item in os.listdir(abs_root):
        old_folder_path = os.path.join(abs_root, item)
        if os.path.isdir(old_folder_path):
            new_item = item
            for old_t, new_t in REPLACEMENTS:
                if old_t in new_item:
                    new_item = new_item.replace(old_t, new_t)
            
            # Rename folder if changed
            if new_item != item:
                new_folder_path = os.path.join(abs_root, new_item)
                if os.path.exists(new_folder_path):
                    shutil.rmtree(new_folder_path)
                os.rename(old_folder_path, new_folder_path)
                print(f"Renamed folder: {item} -> {new_item}")
                target_folder = new_folder_path
            else:
                target_folder = old_folder_path
                
            # Rename .txt files inside
            for f in os.listdir(target_folder):
                if f.endswith('.txt'):
                    new_f = f
                    for old_t, new_t in REPLACEMENTS:
                        if old_t in new_f:
                            new_f = new_f.replace(old_t, new_t)
                    old_f_path = os.path.join(target_folder, f)
                    new_f_path = os.path.join(target_folder, new_f)
                    
                    # Update content of txt
                    with open(old_f_path, 'r', encoding='utf-8', errors='ignore') as fp:
                        content = fp.read()
                    
                    new_content = content
                    for old_t, new_t in REPLACEMENTS:
                        new_content = new_content.replace(old_t, new_t)
                        
                    with open(old_f_path, 'w', encoding='utf-8') as fp:
                        fp.write(new_content)
                        
                    if new_f != f:
                        if os.path.exists(new_f_path):
                            os.remove(new_f_path)
                        os.rename(old_f_path, new_f_path)
                        print(f"Renamed and updated txt: {f} -> {new_f}")

print("\nALL TASKS COMPLETED!")
