import openpyxl
import shutil
import os

updates = {
    '1.2': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email terdaftar.\n3. Isi field Password dengan password yang valid.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Berhasil melakukan login dan diarahkan ke halaman dashboard."
    },
    '1.3': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email yang tidak terdaftar.\n3. Isi field Password dengan password bebas.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error email tidak terdaftar."
    },
    '1.4': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email terdaftar.\n3. Isi field Password dengan password yang salah.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error password salah."
    },
    '1.5': {
        'steps': "1. Pilih Type.\n2. Kosongkan field Email.\n3. Isi field Password dengan password yang valid.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengosongkan Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error field mandatory."
    },
    '1.6': {
        'steps': "1. Pilih Type.\n2. Isi field Email dengan email terdaftar.\n3. Kosongkan field Password.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengosongkan Password.\n4. Menampilkan pesan error field mandatory."
    },
    '1.7': {
        'steps': "1. Pilih Type.\n2. Kosongkan field Email.\n3. Kosongkan field Password.\n4. Klik button LOGIN.",
        'expected': "1. Berhasil memilih Type.\n2. Berhasil mengosongkan Email.\n3. Berhasil mengosongkan Password.\n4. Menampilkan pesan error field mandatory."
    }
}

def update_excel(file_path):
    print(f"Updating {file_path}")
    wb = openpyxl.load_workbook(file_path)
    if 'TC BTN SMART Mobile' not in wb.sheetnames:
        print("  Sheet 'TC BTN SMART Mobile' not found.")
        return
    ws = wb['TC BTN SMART Mobile']
    
    # find column indices
    header_row = 1
    col_tc = col_steps = col_expected = -1
    for i, cell in enumerate(ws[header_row]):
        val = str(cell.value).strip().lower()
        if val == 'no':
            col_tc = i + 1
        elif 'langkah-langkah' in val or 'steps' in val:
            col_steps = i + 1
        elif 'hasil yang diharapkan' in val or 'expected' in val:
            col_expected = i + 1
            
    if col_tc == -1 or col_steps == -1 or col_expected == -1:
        print("  Headers not found!")
        return
        
    updated = 0
    for row in range(2, ws.max_row + 1):
        tc_val = str(ws.cell(row, col_tc).value).strip()
        if tc_val in updates:
            ws.cell(row, col_steps).value = updates[tc_val]['steps']
            ws.cell(row, col_expected).value = updates[tc_val]['expected']
            updated += 1
            
    if updated > 0:
        wb.save(file_path)
        print(f"  Updated {updated} rows and saved.")
    else:
        print("  No rows updated.")

excel_files = [
    r'H:\My Drive\Zegen\BTN Smart\Refactor\Test Script\Test Script BTN Smart Refactor.xlsx',
    r'D:\Project\BTN Smart\Refactor\Test Script\Uji Sistem.xlsx'
]

for ef in excel_files:
    if os.path.exists(ef):
        update_excel(ef)
        # Sync to D and G
        base_name = os.path.basename(ef)
        for drive in ['D', 'G']:
            dst = rf"{drive}:\Project\BTN Smart\Refactor\Skenario\{base_name}" if drive == 'D' else rf"G:\My Drive\Zegen\BTN Smart\Refactor\Skenario\{base_name}"
            if os.path.exists(os.path.dirname(dst)):
                shutil.copy2(ef, dst)
                print(f"  Synced to {dst}")
