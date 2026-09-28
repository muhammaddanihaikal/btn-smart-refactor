import docx, shutil

def remove_row(table, row_idx):
    tr = table.rows[row_idx]._tr
    tr.getparent().remove(tr)

def update_sit(doc_path):
    print(f"Updating SIT document: {doc_path}")
    doc = docx.Document(doc_path)
    
    # 1. Modul 09 (Table 10)
    t9 = doc.tables[10]
    print(f"Table 10 before: {len(t9.rows)} rows. Row 2 is: {t9.rows[2].cells[0].text} - {t9.rows[2].cells[1].text}")
    remove_row(t9, 2) # remove row 2 (9.2)
    # Renumber
    for r_idx in range(1, len(t9.rows)):
        t9.rows[r_idx].cells[0].text = f"9.{r_idx}"
    print(f"Table 10 after: {len(t9.rows)} rows. New numbering:")
    for r_idx in range(1, len(t9.rows)):
        print(f"  {t9.rows[r_idx].cells[0].text} | {t9.rows[r_idx].cells[1].text}")
        
    # 2. Modul 14 (Table 15)
    t14 = doc.tables[15]
    print(f"\nTable 15 before: {len(t14.rows)} rows. Row 2 is: {t14.rows[2].cells[0].text} - {t14.rows[2].cells[1].text}")
    remove_row(t14, 2) # remove row 2 (14.2)
    for r_idx in range(1, len(t14.rows)):
        t14.rows[r_idx].cells[0].text = f"14.{r_idx}"
    print(f"Table 15 after: {len(t14.rows)} rows. New numbering:")
    for r_idx in range(1, len(t14.rows)):
        print(f"  {t14.rows[r_idx].cells[0].text} | {t14.rows[r_idx].cells[1].text}")
        
    # 3. Modul 15 (Table 16)
    t15 = doc.tables[16]
    print(f"\nTable 16 before: {len(t15.rows)} rows. Row 2 is: {t15.rows[2].cells[0].text} - {t15.rows[2].cells[1].text}")
    remove_row(t15, 2) # remove row 2 (15.2)
    for r_idx in range(1, len(t15.rows)):
        t15.rows[r_idx].cells[0].text = f"15.{r_idx}"
    print(f"Table 16 after: {len(t15.rows)} rows. New numbering:")
    for r_idx in range(1, len(t15.rows)):
        print(f"  {t15.rows[r_idx].cells[0].text} | {t15.rows[r_idx].cells[1].text}")

    doc.save(doc_path)
    print(f"Saved: {doc_path}")

local_sit = r'D:\Project\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile.docx'
update_sit(local_sit)

# Sync to Drive H and Drive G
for dst in [
    r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile.docx',
    r'G:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Mobile.docx'
]:
    shutil.copy2(local_sit, dst)
    print(f"Synced SIT to: {dst}")
