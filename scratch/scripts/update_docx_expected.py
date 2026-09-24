import docx
from docx.shared import Pt

doc_path = r'H:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx'
doc = docx.Document(doc_path)

updates = {
    '1.2': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Berhasil melakukan login dan diarahkan ke halaman dashboard.",
    '1.3': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error email tidak terdaftar.",
    '1.4': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error password salah.",
    '1.5': "1. Berhasil memilih Type.\n2. Berhasil mengosongkan Email.\n3. Berhasil mengisi Password.\n4. Menampilkan pesan error field mandatory.",
    '1.6': "1. Berhasil memilih Type.\n2. Berhasil mengisi Email.\n3. Berhasil mengosongkan Password.\n4. Menampilkan pesan error field mandatory.",
    '1.7': "1. Berhasil memilih Type.\n2. Berhasil mengosongkan Email.\n3. Berhasil mengosongkan Password.\n4. Menampilkan pesan error field mandatory."
}

updated_count = 0

for tbl in doc.tables:
    try:
        tc_id = tbl.rows[0].cells[0].text.strip()
        if tc_id in updates:
            expected_text = updates[tc_id]
            # Expected is in Row 2 Cell 1
            # Wait, Table 1.2 had 3 rows (Row 0, 1, 2)
            # R2C1 was: "Hasil yang diharapkan [STATUS]:\n..."
            
            exp_cell = tbl.rows[2].cells[1]
            
            # Preserve [STATUS] if any.
            # In our text, it's just expected. Let's see if [Success] is there.
            old_text = exp_cell.text
            status_suffix = ""
            if "[Success]" in old_text:
                status_suffix = " - [Success]"
            
            # Clear cell
            for p in list(exp_cell.paragraphs):
                p._element.getparent().remove(p._element)
            
            # Recreate
            p = exp_cell.add_paragraph()
            p.paragraph_format.space_before = Pt(5.8)
            run1 = p.add_run("Hasil yang diharapkan [STATUS]:\n")
            run1.font.name = 'Arial'
            run1.font.size = Pt(10)
            run1.bold = True
            
            run2 = p.add_run(expected_text + status_suffix)
            run2.font.name = 'Arial'
            run2.font.size = Pt(10)
            
            if status_suffix:
                # If there's success, we might want to color it, but keeping it simple for now.
                pass
            
            print(f"Updated TC {tc_id} in Word.")
            updated_count += 1
    except Exception as e:
        pass

if updated_count > 0:
    doc.save(doc_path)
    import shutil
    shutil.copy2(doc_path, r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx')
    shutil.copy2(doc_path, r'G:\My Drive\Zegen\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Mobile.docx')
    print("Document saved and synced!")
else:
    print("No tables were updated!")
