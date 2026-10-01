import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc_path = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Web.docx'
doc = docx.Document(doc_path)
base_ss = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Screenshot\Web\33. Menu Absent - Daily Absent'

# Helper to find folder matching tc_prefix
folders = [d for d in os.listdir(base_ss) if os.path.isdir(os.path.join(base_ss, d))]

def get_folder_images(tc_prefix):
    matched = [d for d in folders if d.startswith(tc_prefix + ' ')]
    if not matched:
        return []
    fld = os.path.join(base_ss, matched[0])
    imgs = sorted([f for f in os.listdir(fld) if f.lower().endswith(('.png', '.jpg', '.jpeg'))],
                  key=lambda x: [int(c) for c in os.path.splitext(x)[0].split('.') if c.isdigit()] if any(c.isdigit() for c in os.path.splitext(x)[0]) else [x])
    return [os.path.join(fld, img) for img in imgs]

# 1. Embed Table 68 (33.1)
t68 = doc.tables[68]
c68 = t68.rows[2].cells[1]
c68.text = ''
p68 = c68.paragraphs[0]
p68.alignment = WD_ALIGN_PARAGRAPH.CENTER
imgs_33_1 = get_folder_images('33.1')
for i, img_path in enumerate(imgs_33_1):
    p = p68 if i == 0 else c68.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(img_path, width=Inches(5.9))
print(f'Embedded 33.1: {len(imgs_33_1)} images')

# 2. Embed Table 69 (33.2 to 33.16)
t69 = doc.tables[69]
for k in range(2, 17):
    tc_prefix = f'33.{k}'
    img_row = (k - 2) * 3 + 1
    c = t69.rows[img_row].cells[1]
    c.text = ''
    p0 = c.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    imgs = get_folder_images(tc_prefix)
    for i, img_path in enumerate(imgs):
        p = p0 if i == 0 else c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(img_path, width=Inches(5.9))
    print(f'Embedded {tc_prefix}: {len(imgs)} images')

doc.save(doc_path)
print('Successfully saved Dokumen_Hasil_Uji_Web.docx with all Modul 33 images!')
