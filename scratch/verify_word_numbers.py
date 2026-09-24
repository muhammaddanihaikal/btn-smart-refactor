import docx

doc = docx.Document(r'd:\Project\BTN\SIT\Document SIT BTN Smart\SIT BTN SMART Web.docx')
print(f"Total tables: {len(doc.tables)}")

# Table 0 is header, Table 1 is info
# Tables 2..76 are the 75 section tables
mismatches = []
for sec_idx in range(1, 76):
    tbl_idx = sec_idx + 1
    tbl = doc.tables[tbl_idx]
    rows = tbl.rows
    for r_idx in range(1, len(rows)):
        tc_num = rows[r_idx].cells[0].text.strip()
        expected_num = f"{sec_idx}.{r_idx}"
        if tc_num != expected_num:
            mismatches.append((sec_idx, r_idx, tc_num, expected_num))

print(f"Total number mismatches across all 75 sections: {len(mismatches)}")
if mismatches:
    for m in mismatches[:10]:
        print(f"  Sec {m[0]} row {m[1]}: got '{m[2]}', expected '{m[3]}'")
else:
    print("ALL 75 sections in Word SIT have 100% matching numbering (sec_idx.tc_idx)!")
