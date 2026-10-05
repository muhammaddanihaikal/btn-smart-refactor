import docx
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

doc_p = r'D:\Project\BTN Smart\Refactor\Hasil Uji\Dokumen_Hasil_Uji_Web.docx'
doc = docx.Document(doc_p)

# Canonical clean descriptions for Modul 32 to 42
clean_texts = {
    # Modul 32: Dashboard Absent
    '32.1': 'Berhasil menampilkan halaman Dashboard Absent.',
    '32.2': 'Berhasil menampilkan data dashboard sesuai filter yang dipilih.',
    '32.3': 'Berhasil mereset seluruh filter ke nilai default.',
    '32.4': 'Berhasil menampilkan modal detail daftar absensi Kantor Wilayah.',
    '32.5': 'Berhasil data absensi ditampilkan sesuai kata kunci yang dicari.',
    '32.6': 'Berhasil menampilkan tabel kosong atau pesan data tidak ditemukan.',

    # Modul 33: Daily Absent
    '33.1': 'Berhasil menampilkan halaman Daily Absent.',
    '33.2': 'Berhasil menampilkan data berdasarkan status absent yang dipilih.',
    '33.3': 'Berhasil menampilkan data Daily Absent sesuai filter yang dipilih.',
    '33.4': 'Berhasil mereset filter dan menampilkan data default.',
    '33.5': 'Berhasil membuka drawer Settings.',
    '33.6': 'Berhasil menampilkan data Settings sesuai keyword.',
    '33.7': 'Berhasil menampilkan informasi data tidak ditemukan pada Settings.',
    '33.8': 'Berhasil mengubah status Requires Approval.',
    '33.9': 'Berhasil mengubah status absensi karyawan.',
    '33.10': 'Berhasil menampilkan Attendance Detail.',
    '33.11': 'Berhasil menampilkan Note Clock In.',
    '33.12': 'Berhasil menampilkan Attachment Clock In.',
    '33.13': 'Berhasil menampilkan lokasi Clock In pada peta.',
    '33.14': 'Berhasil menampilkan Note Clock Out.',
    '33.15': 'Berhasil menampilkan Attachment Clock Out.',
    '33.16': 'Berhasil menampilkan lokasi Clock Out pada peta.',

    # Modul 34: Approval Absent
    '34.1': 'Berhasil menampilkan halaman Approval Absent.',
    '34.2': 'Berhasil menampilkan data Approval Absent sesuai filter yang dipilih.',
    '34.3': 'Berhasil mereset filter Approval Absent.',
    '34.4': 'Berhasil mencari data Approval Absent sesuai keyword.',
    '34.5': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '34.6': 'Berhasil menampilkan Note absent.',
    '34.7': 'Berhasil menampilkan Attachment absent.',
    '34.8': 'Berhasil menampilkan lokasi absent pada peta.',
    '34.9': 'Berhasil menyetujui (approve) data absent.',
    '34.10': 'Berhasil menampilkan validasi Comments wajib diisi.',
    '34.11': 'Berhasil menolak (reject) data absent.',
    '34.12': 'Berhasil menampilkan validasi Comments wajib diisi.',
    '34.13': 'Berhasil membatalkan approval yang sudah disetujui.',
    '34.14': 'Berhasil membatalkan approval yang sudah di-reject.',
    '34.15': 'Berhasil mengekspor data approval absent.',

    # Modul 35: Rekap Absent
    '35.1': 'Berhasil menampilkan halaman Rekap Absent.',
    '35.2': 'Berhasil menampilkan daftar data Rekap Absent di dalam tabel.',
    '35.3': 'Berhasil menampilkan data Rekap Absent sesuai filter yang dipilih.',
    '35.4': 'Berhasil mereset filter Rekap Absent ke nilai default.',
    '35.5': 'Berhasil mencari data Rekap Absent sesuai keyword.',
    '35.6': 'Berhasil mendownload file export data Rekap Absent.',

    # Modul 40: Attendance Spot
    '40.1': 'Berhasil menampilkan halaman Attendance Spot.',
    '40.2': 'Berhasil menampilkan data sesuai keyword.',
    '40.3': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '40.4': 'Berhasil menyimpan data Attendance Spot.',
    '40.5': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '40.6': 'Berhasil peta berhasil menampilkan titik lokasi yang dipilih.',
    '40.7': 'Peta tidak menemukan lokasi yang dimaksud (menampilkan pesan tidak ditemukan).',
    '40.8': 'Berhasil menyimpan perubahan data Attendance Spot.',
    '40.9': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '40.10': 'Berhasil menampilkan popup peta lokasi Attendance Spot.',
    '40.11': 'Berhasil menghapus data Attendance Spot.',
    '40.12': 'Berhasil menampilkan halaman Setting Attendance Spot Personnel.',
    '40.13': 'Berhasil menampilkan data Personnel sesuai filter yang dipilih.',
    '40.14': 'Berhasil menampilkan data Personnel sesuai keyword.',
    '40.15': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '40.16': 'Berhasil menambahkan/mengaktifkan personnel ke Attendance Spot secara single.',
    '40.17': 'Berhasil menghapus/menonaktifkan personnel dari Attendance Spot secara single.',
    '40.18': 'Berhasil menambahkan/mengaktifkan seluruh personnel yang dipilih ke Attendance Spot secara bulk.',
    '40.19': 'Berhasil menghapus/menonaktifkan seluruh personnel yang dipilih dari Attendance Spot secara bulk.',

    # Modul 41: Work Pattern
    '41.1': 'Berhasil menampilkan halaman Work Pattern.',
    '41.2': 'Berhasil menampilkan data sesuai keyword.',
    '41.3': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '41.4': 'Berhasil menyimpan data Work Pattern.',
    '41.5': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '41.6': 'Berhasil menerapkan jadwal yang diinput ke seluruh hari yang dicentang secara bersamaan.',
    '41.7': 'Berhasil mengosongkan kembali (reset) jadwal pada seluruh hari yang dicentang secara bersamaan.',
    '41.8': 'Berhasil menyimpan perubahan data Work Pattern.',
    '41.9': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '41.10': 'Berhasil menampilkan popup Work Pattern Details yang berisi jadwal kerja.',
    '41.11': 'Berhasil menghapus data Work Pattern.',
    '41.12': 'Berhasil menampilkan halaman Setting Work Pattern Personnel.',
    '41.13': 'Berhasil menampilkan data Personnel sesuai filter yang dipilih.',
    '41.14': 'Berhasil menampilkan data Personnel sesuai keyword.',
    '41.15': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '41.16': 'Berhasil meng-assign personnel tersebut ke Work Pattern.',
    '41.17': 'Berhasil meng-assign seluruh personnel yang dipilih secara bulk ke Work Pattern.',
    '41.18': 'Berhasil menghapus (unassign) personnel tersebut dari Work Pattern.',
    '41.19': 'Berhasil menghapus (unassign) seluruh personnel yang dipilih secara bulk dari Work Pattern.',

    # Modul 42: Holiday
    '42.1': 'Berhasil menampilkan halaman Holiday.',
    '42.2': 'Berhasil menampilkan data sesuai keyword.',
    '42.3': 'Berhasil menampilkan informasi data tidak ditemukan.',
    '42.4': 'Berhasil menyimpan data Holiday.',
    '42.5': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '42.6': 'Berhasil menyimpan perubahan data Holiday.',
    '42.7': 'Berhasil menampilkan validasi field mandatory wajib diisi.',
    '42.8': 'Berhasil menghapus data Holiday.',
}

def set_clean_expected_result(cell, desc_text):
    # Strip any '[Success]' or number prefix like '1. ', '2. '
    clean_desc = desc_text.strip()
    clean_desc = re.sub(r'^\d+\.\s*', '', clean_desc)
    clean_desc = re.sub(r'\s*-\s*\[Success\]', '', clean_desc).strip()
    
    # 1. Clear cell text (resets all paragraphs to 1 empty paragraph)
    cell.text = ''
    
    # Paragraph 0: Header
    p0 = cell.paragraphs[0]
    p0_pr = p0._p.get_or_add_pPr()
    p0_sp = OxmlElement('w:spacing')
    p0_sp.set(qn('w:before'), '116')
    p0_sp.set(qn('w:after'), '0')
    p0_pr.append(p0_sp)
    
    r0 = p0.add_run('Hasil yang diharapkan [STATUS]:')
    r0.font.name = 'Arial'
    r0.font.size = Pt(10)
    r0.bold = True
    
    # Paragraph 1: Description + [Success]
    p1 = cell.add_paragraph()
    p1_pr = p1._p.get_or_add_pPr()
    p1_sp = OxmlElement('w:spacing')
    p1_sp.set(qn('w:before'), '60')
    p1_sp.set(qn('w:after'), '100')
    p1_pr.append(p1_sp)
    
    # Run 1: Description in #2C5293
    r1 = p1.add_run(clean_desc)
    r1.font.name = 'Arial'
    r1.font.size = Pt(10)
    r1_pr = r1._element.get_or_add_rPr()
    c1 = OxmlElement('w:color')
    c1.set(qn('w:val'), '2C5293')
    r1_pr.append(c1)
    
    # Run 2: ' - [Success]' in bold #2C5293
    r2 = p1.add_run(' - [Success]')
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.bold = True
    r2_pr = r2._element.get_or_add_rPr()
    c2 = OxmlElement('w:color')
    c2.set(qn('w:val'), '2C5293')
    r2_pr.append(c2)

# Apply to all relevant modules
current_h = ''
updated_count = 0

for el in doc.element.body:
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(el, doc)
        if p.style.name.startswith('Heading 1'):
            current_h = p.text
    elif tag == 'tbl':
        for m in ['32', '33', '34', '35', '40', '41', '42']:
            if f'{m}. Modul' in current_h:
                tbl = docx.table.Table(el, doc)
                # Check Table 0 (first TC .1)
                if len(tbl.rows) == 4 and tbl.rows[1].cells[0].text.strip() == f"{m}.1":
                    tc_key = f"{m}.1"
                    if tc_key in clean_texts:
                        set_clean_expected_result(tbl.rows[3].cells[1], clean_texts[tc_key])
                        updated_count += 1
                        print(f"Updated {tc_key} in Table 0: {clean_texts[tc_key]}")
                # Check Table 1 (.2 onwards)
                else:
                    for r_idx in range(0, len(tbl.rows), 3):
                        if r_idx < len(tbl.rows):
                            tc_val = tbl.rows[r_idx].cells[0].text.strip()
                            r_exp = r_idx + 2
                            if r_exp < len(tbl.rows) and tc_val in clean_texts:
                                set_clean_expected_result(tbl.rows[r_exp].cells[1], clean_texts[tc_val])
                                updated_count += 1
                                print(f"Updated {tc_val} in Table 1 (row {r_exp}): {clean_texts[tc_val]}")

doc.save(doc_p)
print(f"\nSuccessfully cleaned up and standardized {updated_count} expected result cells in {doc_p}!")
