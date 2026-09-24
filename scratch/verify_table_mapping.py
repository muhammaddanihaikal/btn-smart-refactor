import docx
import openpyxl
import re
from collections import Counter

EXCEL_PATH = r'd:\Project\BTN\SIT\Test Case.xlsx'
DOC_PATH = r'H:\My Drive\Zegen\BTN Smart\Refactor\SIT\SIT BTN SMART Web (Updated).docx'

# Load Excel mapping
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb['TC BTN SMART Web']

excel_map = {}
for r in range(2, ws.max_row + 1):
    m = ws.cell(r, 2).value
    s = ws.cell(r, 3).value
    t = ws.cell(r, 4).value
    if m and t:
        m_str = str(m).strip()
        s_str = str(s).strip() if s else ''
        t_clean = str(t).strip().lower()
        excel_map[(m_str, t_clean)] = s_str

doc = docx.Document(DOC_PATH)
body = doc.element.body

modules_list = [
    'Login', 'Profile Nasabah & Sales', 'Profile', 'Overview', 'User Authority', 
    'Bisnis dan Produk', 'Kantor', 'Menu Target', 'Upload Bulk', 
    'Lead Generation', 'Lead Qualification', 'List Prospek ETB', 
    'Menu Absent', 'Sales Tracking Activity', 'Setting Absent', 
    'Sales Force', 'Report Funding', 'Report Lending', 
    'Re-Assign dan Approval', 'Export data management'
]

cur_mod = 'Unknown'
tables_mapped = []

for child in list(body):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        t = ''.join(child.itertext()).strip()
        if t and 'Modul' in t:
            for m in modules_list:
                if m in t:
                    cur_mod = m
                    break
    elif tag == 'tbl':
        rows = child.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
        if rows:
            cells = rows[0].findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
            if len(cells) == 7 and len(rows) > 1:
                # Deduce sub menu from rows
                sub_votes = []
                for r in rows[1:]:
                    r_cells = r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
                    title = ''.join(r_cells[1].itertext()).strip().lower()
                    if (cur_mod, title) in excel_map:
                        sub_votes.append(excel_map[(cur_mod, title)])
                    else:
                        for (em, et), es in excel_map.items():
                            if em == cur_mod and (et in title or title in et):
                                sub_votes.append(es)
                                break
                sub_name = Counter(sub_votes).most_common(1)[0][0] if sub_votes else ''
                tables_mapped.append((cur_mod, sub_name, len(rows) - 1))

print(f"Total tables mapped: {len(tables_mapped)}")
for idx, (m, s, count) in enumerate(tables_mapped, 1):
    label = f"Modul {m} - {s}" if s else f"Modul {m}"
    print(f"Table {idx:2d} ({count:2d} TCs) -> {label}")
